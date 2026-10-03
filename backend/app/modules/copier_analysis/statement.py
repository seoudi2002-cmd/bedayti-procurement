"""Readers for the monthly copier consumption statement.

* Part 1 (the primary operational source): blocks per machine class/package, one row per machine
  (branch, previous reading, current reading, consumption, additional copies).
* Part 2 (supporting detail, in the same workbook) and the Word statement: the same machines grouped by governorate,
  with machine class and rental value. Never an independent total: it is matched to Part 1 machine by machine.

Nothing is corrected: disagreements become issues. Blank is not zero. The source file is never modified."""
import io
import re
from dataclasses import dataclass, field

from openpyxl import load_workbook

from app.core.cleaning.normalizers import normalize_text

PACKAGES = (1000, 2000, 3000, 5000, 10000)
_STRAY_B = re.compile(r"(?<![A-Za-z])B(?![A-Za-z])")


class UnrecognisedStatement(ValueError):
    pass


@dataclass
class Machine:
    source_ref: str
    class_key: str
    package: int | None
    branch_source: str
    branch_key: str
    seq: int | None
    prev: int | None
    cur: int | None
    cons: int | None
    exc: int | None
    flags: list[str] = field(default_factory=list)
    # supporting detail (Part 2 / Word), filled by matching
    location_group: str | None = None
    group_kind: str | None = None
    rent_detail: int | None = None
    detail_sources: list[str] = field(default_factory=list)


@dataclass
class DetailRow:
    source_ref: str
    source_kind: str  # xlsx_part2 | word
    location_group: str | None
    group_kind: str | None
    branch_source: str
    class_key: str | None
    machine_text: str
    prev: int | None
    cur: int | None
    cons: int | None
    exc: int | None
    rent: int | None


@dataclass
class Block:
    class_key: str
    title: str
    package: int | None
    machines: int
    stated_excess_total: int | None
    computed_excess: int
    total_row_ref: str | None


@dataclass
class Issue:
    code: str
    severity: str
    message: str
    count: int = 1
    examples: list[str] = field(default_factory=list)


@dataclass
class Statement:
    period: tuple[int, int] | None = None  # (year, month) as written in the title
    title: str | None = None
    machines: list[Machine] = field(default_factory=list)
    blocks: list[Block] = field(default_factory=list)
    details: list[DetailRow] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)

    def issue(self, code: str, severity: str, message: str, example: str | None = None) -> None:
        for i in self.issues:
            if i.code == code:
                i.count += 1
                if example and len(i.examples) < 10:
                    i.examples.append(example)
                return
        self.issues.append(Issue(code, severity, message, 1, [example] if example else []))


# ---------------------------------------------------------------------------------------------- helpers
def num(v) -> int | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else None
    s = re.sub(r"[,\s ]", "", str(v))
    return int(s) if re.fullmatch(r"-?\d+", s) else None


def clean_label(s) -> str:
    return " ".join(str(s or "").replace(" ", " ").split())


def display_branch(label: str) -> str:
    """Presentation only (the source text is stored untouched): drop a stray Latin 'B' marker and extra dashes/spaces."""
    t = _STRAY_B.sub(" ", clean_label(label))
    t = re.sub(r"(?:^|\s)-+(?:\s|$)", " ", t).strip(" -")
    t = re.sub(r"(?<=[؀-ۿ])B$", "", t).strip()
    return " ".join(t.split())


def branch_key(label: str) -> str:
    t = re.sub(r"(?<=[؀-ۿ])B$", "", clean_label(label))
    return normalize_text(_STRAY_B.sub(" ", t)).replace("-", " ").strip()


def has_stray_marker(label: str) -> bool:
    t = clean_label(label)
    return bool(_STRAY_B.search(t)) or bool(re.search(r"(?<=[؀-ۿ])B$", t))


def classify(text: str) -> tuple[str | None, int | None]:
    """class_key and package from a block title or a machine-type text; (None, None) when it cannot be told."""
    n = normalize_text(text)
    pkgs = [int(x) for x in re.findall(r"\d+", n) if int(x) in PACKAGES]
    if not pkgs:
        return None, None
    pkg = max(pkgs)
    color = "الوان" in n
    if color and ("زيروكس" in n or "a3" in n):
        return f"color_a3_{pkg}", pkg
    if color:
        return f"color_copier_{pkg}", pkg
    if "طابعات" in n or "طابعه" in n:
        return f"printer_{pkg}", pkg
    return f"mono_copier_{pkg}", pkg


def _period(text: str) -> tuple[int, int] | None:
    m = re.search(r"(?<!\d)(\d{1,2})\s*-\s*(20\d{2})(?!\d)", text)
    return (int(m.group(2)), int(m.group(1))) if m and 1 <= int(m.group(1)) <= 12 else None


def _header_columns(ws, r: int) -> dict[str, int]:
    cols: dict[str, int] = {}
    for c in range(1, 16):
        n = normalize_text(ws.cell(r, c).value)
        if not n:
            continue
        for key, words in (("branch", ("الفرع",)), ("machine", ("الماكينه",)), ("prev", ("السابقه",)), ("cur", ("الحاليه",)),
                           ("cons", ("الاستهلاك",)), ("exc", ("الاضافي", "الاضافى")), ("rent", ("الايجاريه",)), ("notes", ("ملاحظات",)),
                           ("seq", ("م",))):
            if (key == "seq" and n == "م") or (key != "seq" and any(w in n for w in words)):
                cols.setdefault(key, c)
    return cols


def _is_total_label(v) -> bool:
    n = normalize_text(v)
    return n.startswith("الاجمالي") or n == "اجمالي" or n == "total"


# ---------------------------------------------------------------------------------------------- xlsx
def parse_statement_xlsx(content: bytes) -> Statement:
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = next((w for w in wb.worksheets if any("زيارات" in str(c.value or "") for row in w.iter_rows(max_row=60) for c in row)), None)
    if ws is None:
        raise UnrecognisedStatement("No consumption-statement sheet found (no 'زيارات ... ماكينات' block titles)")
    st = Statement()
    last = max((r for (r, _c), cell in ws._cells.items() if cell.value is not None), default=0)
    r = 1
    cols: dict[str, int] = {}
    block: dict | None = None
    group, group_kind = None, None
    part2 = False

    def close_block():
        nonlocal block
        if block:
            st.blocks.append(Block(block["class_key"], block["title"], block["package"], block["machines"], block["stated"],
                                   block["computed"], block["total_ref"]))
            if block["stated"] is None:
                st.issue("block_total_missing", "warning", "Block without its own 'total' row (additional copies not stated)",
                         f"{ws.title}!{block['title'][:40]}")
            elif block["stated"] != block["computed"]:
                st.issue("block_total_differs", "warning", "Block's stated additional-copies total differs from its rows",
                         f"{block['total_ref']}: stated {block['stated']} vs rows {block['computed']}")
        block = None

    while r <= last:
        texts = {c: ws.cell(r, c).value for c in range(1, 16) if ws.cell(r, c).value is not None}
        if not texts:
            r += 1
            continue
        title = next((str(v) for v in texts.values() if isinstance(v, str) and "زيارات" in v), None)
        if title:  # Part 1 block title
            close_block()
            part2 = False
            st.title = st.title or clean_label(title)
            st.period = st.period or _period(title)
            ck, pk = classify(title)
            if ck is None:
                st.issue("block_class_unreadable", "critical", "Block title does not state a known package size", f"{ws.title}!R{r}")
            block = {"class_key": ck or "unknown", "title": clean_label(title), "package": pk, "machines": 0, "stated": None,
                     "computed": 0, "total_ref": None}
            cols = {}
            r += 1
            continue
        hdr = _header_columns(ws, r)
        if "branch" in hdr and "prev" in hdr and "cur" in hdr:  # a table header row
            cols = hdr
            if "machine" in hdr:  # Part 2 header: a new supporting section
                close_block()
                part2 = True
            r += 1
            continue
        d = ws.cell(r, cols.get("branch", 4)).value if cols else None
        if cols and any(_is_total_label(v) for v in texts.values() if isinstance(v, str)) and block and not part2:
            ev = ws.cell(r, cols["exc"]).value if "exc" in cols else None
            block["stated"], block["total_ref"] = num(ev), f"{ws.title}!R{r}"
            r += 1
            continue
        if cols and "branch" in cols:
            prev, cur = num(ws.cell(r, cols["prev"]).value), num(ws.cell(r, cols["cur"]).value)
            if isinstance(d, str) and d.strip() and prev is not None and cur is not None:
                ref = f"{ws.title}!R{r}"
                cons = num(ws.cell(r, cols["cons"]).value) if "cons" in cols else None
                exc = num(ws.cell(r, cols["exc"]).value) if "exc" in cols else None
                if part2:
                    mt = clean_label(ws.cell(r, cols["machine"]).value)
                    ck, _pk = classify(mt) if not mt.isdigit() else classify(mt)
                    st.details.append(DetailRow(ref, "xlsx_part2", group, group_kind, clean_label(d), ck, mt, prev, cur, cons, exc,
                                                num(ws.cell(r, cols["rent"]).value) if "rent" in cols else None))
                elif block is not None:
                    seq = num(ws.cell(r, cols["seq"]).value) if "seq" in cols else None
                    m = Machine(ref, block["class_key"], block["package"], clean_label(d), branch_key(d), seq, prev, cur, cons, exc)
                    st.machines.append(m)
                    block["machines"] += 1
                    block["computed"] += exc or 0
                r += 1
                continue
            # section title row of Part 2 (a governorate / administrations name) = text only in the branch column
            if isinstance(d, str) and d.strip() and len(texts) <= 2 and not any(isinstance(v, (int, float)) for v in texts.values()):
                pass
        # governorate / section title (text in one cell, nothing numeric)
        only = [v for v in texts.values() if isinstance(v, str)]
        if len(only) == 1 and len(texts) == 1 and not title:
            group = clean_label(only[0])
            group_kind = "administrations" if normalize_text(group).startswith("الادار") else "governorate"
            part2 = part2 or False
        r += 1
    close_block()
    if not st.machines:
        raise UnrecognisedStatement("No machine rows found in the statement")
    return st


# ---------------------------------------------------------------------------------------------- Word
def parse_statement_word(content: bytes, filename: str = "") -> list[DetailRow]:
    """Detail rows from the Word statement (.doc legacy binary or .docx). Supporting source only."""
    if filename.lower().endswith(".docx") or content[:2] == b"PK":
        return _word_docx(content)
    from app.core.extraction.doc_legacy import read_doc_text
    return _word_tokens(read_doc_text(content).split("\r"))


def _group_kind(title: str) -> str:
    return "administrations" if normalize_text(title).startswith("الادار") else "governorate"


def _word_tokens(paragraphs: list[str]) -> list[DetailRow]:
    rows: list[DetailRow] = []
    for pi, para in enumerate(p for p in paragraphs if len(p) > 120 and "\x07" in p):
        t = [x.replace(" ", " ").strip() for x in para.split("\x07")]
        heads = [i for i in range(len(t) - 1) if t[i] == "م" and normalize_text(t[i + 1]) == "الفرع"]
        for hi, h in enumerate(heads):
            title = next((t[k] for k in range(h - 1, -1, -1) if t[k]), None)
            tail = h
            while tail < len(t) and t[tail]:
                tail += 1
            labels = t[h:tail]
            stride = len(labels) + 1
            cols = {normalize_text(x): j for j, x in enumerate(labels)}
            idx = {k: next((j for n, j in cols.items() if any(w in n for w in ws)), None) for k, ws in (
                ("branch", ("الفرع",)), ("machine", ("الماكينه",)), ("prev", ("السابقه",)), ("cur", ("الحاليه",)),
                ("cons", ("الاستهلاك",)), ("exc", ("الاضافي", "الاضافى")), ("rent", ("الايجاريه",)))}
            end = len(t)
            if hi + 1 < len(heads):  # stop before the next section's title
                nxt = heads[hi + 1]
                end = next((k for k in range(nxt - 1, h, -1) if t[k]), nxt)
            body = t[tail + 1:end]
            for k in range(0, len(body) - stride + 1, stride):
                cell = body[k:k + stride]
                br = cell[idx["branch"]] if idx["branch"] is not None else ""
                prev, cur = num(cell[idx["prev"]]), num(cell[idx["cur"]])
                if not br or prev is None or cur is None:
                    continue
                mt = cell[idx["machine"]] if idx["machine"] is not None else ""
                ck, _ = classify(mt)
                rows.append(DetailRow(f"word:p{pi + 1}", "word", title, _group_kind(title or ""), clean_label(br), ck, mt, prev, cur,
                                      num(cell[idx["cons"]]) if idx["cons"] is not None else None,
                                      num(cell[idx["exc"]]) if idx["exc"] is not None else None,
                                      num(cell[idx["rent"]]) if idx["rent"] is not None else None))
    return rows


def _word_docx(content: bytes) -> list[DetailRow]:
    from docx import Document
    doc = Document(io.BytesIO(content))
    rows: list[DetailRow] = []
    for ti, table in enumerate(doc.tables):
        grid = [[clean_label(c.text) for c in row.cells] for row in table.rows]
        hi = next((i for i, r in enumerate(grid) if "م" in r and any(normalize_text(x) == "الفرع" for x in r)), None)
        if hi is None:
            continue
        title = next((c for r in grid[:hi] for c in r if c), None)
        head = [normalize_text(x) for x in grid[hi]]

        def col(*ws):
            return next((j for j, n in enumerate(head) if any(w in n for w in ws)), None)
        idx = {"branch": col("الفرع"), "machine": col("الماكينه"), "prev": col("السابقه"), "cur": col("الحاليه"), "cons": col("الاستهلاك"),
               "exc": col("الاضافي", "الاضافى"), "rent": col("الايجاريه")}
        for ri, r in enumerate(grid[hi + 1:], hi + 2):
            br = r[idx["branch"]] if idx["branch"] is not None else ""
            prev, cur = num(r[idx["prev"]]) if idx["prev"] is not None else None, num(r[idx["cur"]]) if idx["cur"] is not None else None
            if not br or prev is None or cur is None:
                continue
            mt = r[idx["machine"]] if idx["machine"] is not None else ""
            rows.append(DetailRow(f"word:t{ti + 1}r{ri}", "word", title, _group_kind(title or ""), br, classify(mt)[0], mt, prev, cur,
                                  num(r[idx["cons"]]) if idx["cons"] is not None else None,
                                  num(r[idx["exc"]]) if idx["exc"] is not None else None,
                                  num(r[idx["rent"]]) if idx["rent"] is not None else None))
    return rows


# ---------------------------------------------------------------------------------------------- validation + matching
def validate_machines(st: Statement) -> None:
    for m in st.machines:
        if m.prev is not None and m.cur is not None:
            if m.cur < m.prev:
                m.flags.append("reading_decreased")
                st.issue("reading_decreased", "critical", "Current reading is lower than the previous one", m.source_ref)
            if m.cons is not None and m.cur - m.prev != m.cons:
                m.flags.append("consumption_mismatch")
                st.issue("consumption_mismatch", "warning", "Consumption differs from current − previous reading", m.source_ref)
        if m.cons is None:
            m.flags.append("consumption_missing")
            st.issue("consumption_missing", "warning", "Machine row without consumption", m.source_ref)
        if m.exc is None:
            m.flags.append("excess_missing")
            st.issue("excess_missing", "warning", "Machine row without a stated additional-copies value", m.source_ref)
        elif m.package and m.cons is not None and m.exc != max(0, m.cons - m.package):
            m.flags.append("excess_differs_from_rule")
            st.issue("excess_differs_from_rule", "warning",
                     "Stated additional copies differ from max(0, consumption − package)", f"{m.source_ref}: {m.exc} vs {max(0, m.cons - m.package)}")
        if m.cons is not None and m.cons == 0:
            m.flags.append("zero_consumption")
        if has_stray_marker(m.branch_source):
            m.flags.append("branch_stray_marker")
            st.issue("branch_stray_marker", "info", "Branch label carries a stray Latin 'B' marker (shown without it; source kept)", m.source_ref)
    seen: dict[tuple, str] = {}
    for m in st.machines:
        k = (m.class_key, m.prev, m.cur)
        if k in seen:
            m.flags.append("duplicate_machine_reading")
            st.issue("duplicate_machine_reading", "warning", "Two machines of one class share the same readings", f"{seen[k]} / {m.source_ref}")
        seen[k] = m.source_ref


def _match_class(class_key: str | None) -> str | None:
    """The detail tables state only a package for mono machines (no copier/printer word), so those match by package."""
    if class_key and (class_key.startswith("printer_") or class_key.startswith("mono_copier_")):
        return "mono_" + class_key.rsplit("_", 1)[1]
    return class_key


def match_details(st: Statement, details: list[DetailRow], source_kind: str) -> None:
    """Attach supporting detail (governorate, rental value) to Part 1 machines by class + readings. Disagreements and
    unmatched rows on either side become issues; Part 1 values are never changed."""
    by_key: dict[tuple, list[Machine]] = {}
    for m in st.machines:
        by_key.setdefault((_match_class(m.class_key), m.prev, m.cur), []).append(m)
    used: set[int] = set()
    label = {"xlsx_part2": "Part 2", "word": "the Word statement"}[source_kind]
    for d in details:
        cands = [m for m in by_key.get((_match_class(d.class_key), d.prev, d.cur), []) if id(m) not in used]
        if not cands:
            st.issue(f"detail_row_not_in_part1_{source_kind}", "warning", f"A row of {label} has no matching machine in Part 1", d.source_ref)
            continue
        m = cands[0]
        used.add(id(m))
        m.detail_sources.append(source_kind)
        if source_kind == "xlsx_part2" or m.location_group is None:
            m.location_group, m.group_kind = d.location_group, d.group_kind
        if d.rent is not None and source_kind == "xlsx_part2":
            m.rent_detail = d.rent
        if d.cons is not None and d.cons != m.cons or d.exc is not None and d.exc != m.exc:
            st.issue(f"detail_value_differs_{source_kind}", "warning", f"Consumption/additional copies in {label} differ from Part 1", f"{d.source_ref} vs {m.source_ref}")
        if branch_key(d.branch_source) != m.branch_key:
            st.issue(f"detail_branch_name_differs_{source_kind}", "info", f"Branch spelling in {label} differs from Part 1 (matched by readings)",
                     f"{m.branch_source} ≠ {d.branch_source}")
    missing = [m for m in st.machines if source_kind not in m.detail_sources]
    for m in missing:
        st.issue(f"machine_missing_in_detail_{source_kind}", "warning", f"A Part 1 machine is missing from {label}", f"{m.source_ref} {display_branch(m.branch_source)}")
