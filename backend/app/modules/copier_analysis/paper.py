"""Copier-paper distribution statement: reader + analysis.

The owner defines *distribution = consumption* (cartons given to a branch / head-office department, dated). Cost =
cartons x the carton price of the purchase window (owner-provided schedule, VAT-exempt paper). A carton = 5 reams x 500
sheets (owner-provided). The *need* implied by the machines is a separate estimate: pages / pages-per-sheet / sheets-per-carton.
Source rows are never changed; dubious rows are flagged."""
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal

import yaml
from openpyxl import load_workbook

from app.core.cleaning.normalizers import normalize_text
from app.modules.copier_analysis.statement import branch_key as _strip_branch_key
from app.modules.copier_analysis.statement import clean_label

ZERO = Decimal("0")
_MONTHS = {"يناير": 1, "فبراير": 2, "مارس": 3, "ابريل": 4, "مايو": 5, "يونيو": 6, "يوليو": 7, "اغسطس": 8, "سبتمبر": 9, "اكتوبر": 10, "نوفمبر": 11, "ديسمبر": 12}
_HQ = normalize_text("المركز الرئيسي")


class UnrecognisedPaperStatement(ValueError):
    pass


@dataclass
class PaperRow:
    source_ref: str
    seq: int | None
    cartons: Decimal
    branch_source: str
    branch_display: str
    branch_key: str
    is_head_office: bool
    department: str | None
    distributed_on: date | None
    flags: list[str] = field(default_factory=list)


@dataclass
class PaperIssue:
    code: str
    severity: str
    message: str
    count: int = 1
    examples: list[str] = field(default_factory=list)


@dataclass
class PaperStatement:
    title: str | None = None
    po_no: str | None = None
    receipt_date: date | None = None
    received_cartons: Decimal | None = None
    stated_total: Decimal | None = None
    stated_months: dict[int, Decimal] = field(default_factory=dict)   # month -> the statement's own monthly subtotal
    rows: list[PaperRow] = field(default_factory=list)
    issues: list[PaperIssue] = field(default_factory=list)

    def issue(self, code: str, severity: str, message: str, example: str | None = None) -> None:
        for i in self.issues:
            if i.code == code:
                i.count += 1
                if example and len(i.examples) < 10:
                    i.examples.append(example)
                return
        self.issues.append(PaperIssue(code, severity, message, 1, [example] if example else []))


def _month(text) -> int | None:
    for tok in normalize_text(text).split():
        if tok in _MONTHS:
            return _MONTHS[tok]
    return None


def split_head_office(label: str) -> tuple[bool, str | None]:
    """'المركز الرئيسي - الحفظ' / 'المركز الرئيسي الدور الثاني' -> (True, 'الحفظ' / 'الدور الثاني'); plain 'المركز الرئيسي' -> (True, None)."""
    n = normalize_text(label)
    if not n.startswith(_HQ):
        return False, None
    rest = clean_label(re.sub(r"^\s*المركز\s+الرئيس[يى]\s*[-–—:]*\s*", "", clean_label(label)))
    return True, (rest or None)


def is_distribution_workbook(content: bytes) -> bool:
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(max_row=3, values_only=True):
                if any(isinstance(v, str) and "بيان توزيع ورق" in v for v in row):
                    return True
    except Exception:
        return False
    return False


def parse_distribution_xlsx(content: bytes) -> PaperStatement:
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = next((w for w in wb.worksheets if any("بيان توزيع ورق" in str(c.value or "") for row in w.iter_rows(max_row=3) for c in row)), None)
    if ws is None:
        raise UnrecognisedPaperStatement("No 'بيان توزيع ورق' (paper distribution) sheet found")
    st = PaperStatement()
    title = next(str(c.value) for row in ws.iter_rows(max_row=3) for c in row if "بيان توزيع ورق" in str(c.value or ""))
    st.title = clean_label(title)
    if m := re.search(r"أمر شراء رقم\s*\(?\s*(\d+)\s*\)?", title):
        st.po_no = m.group(1)
    if m := re.search(r"من تاريخ\s*(\d{1,2})\s*(\S+)\s*(20\d{2})", title):
        mo = _month(m.group(2))
        if mo:
            st.receipt_date = date(int(m.group(3)), mo, int(m.group(1)))
    if m := re.search(r"استلام\s*(\d+)\s*كرتونة", title):
        st.received_cartons = Decimal(m.group(1))
    hdr = None
    for r in range(1, 8):
        cols = [(normalize_text(ws.cell(r, c).value), c) for c in range(1, 12) if ws.cell(r, c).value is not None]
        labels = {k: c for k, c in reversed(cols)}  # first occurrence wins
        if "الفرع" in labels and "التاريخ" in labels:
            qtys = [c for k, c in cols if k == "العدد"]
            if not qtys:
                break
            hdr = {"branch": labels["الفرع"], "date": labels["التاريخ"], "seq": labels.get("م", 1), "qty": qtys[0],
                   "sub": qtys[1] if len(qtys) > 1 else None, "row": r}
            break
    if hdr is None:
        raise UnrecognisedPaperStatement("Header row (م / العدد / الفرع / التاريخ) not found")
    qty_col, sub_col = hdr["qty"], hdr["sub"]  # monthly subtotals sit in the 2nd "العدد" column
    last = max((r for (r, _c), cell in ws._cells.items() if cell.value is not None), default=0)
    for r in range(hdr["row"] + 1, last + 1):
        label_a = ws.cell(r, 1).value
        if isinstance(label_a, str) and normalize_text(label_a).startswith("العدد الاجمالي"):
            st.stated_total = _num(ws.cell(r, 2).value)
            continue
        branch = ws.cell(r, hdr["branch"]).value
        qty = _num(ws.cell(r, qty_col).value)
        sub = ws.cell(r, sub_col).value if sub_col else None
        note = ws.cell(r, sub_col + 1).value if sub_col else None
        if isinstance(note, str) and "استهلاك شهر" in note and isinstance(sub, (int, float)):
            mo = _month(note)
            if mo:
                st.stated_months[mo] = Decimal(str(sub))
        if not branch or qty is None:
            continue
        d = ws.cell(r, hdr["date"]).value
        d = d.date() if isinstance(d, datetime) else None
        seq = ws.cell(r, hdr["seq"]).value
        label = clean_label(branch)
        hq, dept = split_head_office(label)
        row = PaperRow(f"{ws.title}!R{r}", int(seq) if isinstance(seq, (int, float)) else None, qty, label, label,
                       normalize_text(_strip_branch_key(label)), hq, dept, d)
        if d is None:
            row.flags.append("date_missing")
            st.issue("date_missing", "critical", "Distribution line without a date", row.source_ref)
        st.rows.append(row)
    if not st.rows:
        raise UnrecognisedPaperStatement("No distribution rows found")
    _validate(st)
    return st


def _num(v) -> Decimal | None:
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    return None


def _validate(st: PaperStatement) -> None:
    total = sum((r.cartons for r in st.rows), ZERO)
    if st.stated_total is not None and st.stated_total != total:
        st.issue("total_differs", "warning", "The statement's stated total differs from the sum of its lines", f"stated {st.stated_total} vs lines {total}")
    if st.received_cartons is not None and st.received_cartons != total:
        st.issue("received_differs", "warning", "Cartons received (title) differ from cartons distributed", f"received {st.received_cartons} vs distributed {total}")
    by_m: dict[int, Decimal] = {}
    for r in st.rows:
        if r.distributed_on:
            by_m[r.distributed_on.month] = by_m.get(r.distributed_on.month, ZERO) + r.cartons
    for mo, stated in st.stated_months.items():
        if by_m.get(mo, ZERO) != stated:
            st.issue("month_subtotal_differs", "warning", "A monthly subtotal stated in the file differs from its lines", f"month {mo}: stated {stated} vs lines {by_m.get(mo, ZERO)}")
    seqs = [(r.seq, r.source_ref) for r in st.rows if r.seq is not None]
    present = {s for s, _ in seqs}
    gaps = [i for i in range(1, max(present) + 1) if i not in present] if present else []
    if gaps:
        st.issue("sequence_gaps", "info", "Missing numbers in the running sequence (م)", "missing: " + ", ".join(map(str, gaps[:12])))
    for (a, _ra), (b, rb) in zip(seqs, seqs[1:]):
        if b < a:
            st.issue("sequence_out_of_order", "info", "Running number is lower than the previous line's", rb)
    seen: dict[tuple, list[PaperRow]] = {}
    for r in st.rows:
        seen.setdefault((r.branch_key, r.distributed_on), []).append(r)
    for (k, d), rs in seen.items():
        if len(rs) > 1 and d:
            for r in rs:
                r.flags.append("same_branch_same_day")
            st.issue("same_branch_same_day", "warning", "Same branch twice on the same day (possible duplicate or a second delivery — kept, review)",
                     f"{rs[0].branch_display} {d.isoformat()}: " + ", ".join(x.source_ref for x in rs))
    variants: dict[str, set[str]] = {}
    for r in st.rows:
        if r.is_head_office:
            variants.setdefault(normalize_text(r.department or ""), set()).add(r.branch_source)
    for k, v in variants.items():
        if len(v) > 1:
            st.issue("head_office_label_variants", "info", "One head-office unit written in several ways (grouped by the text after 'المركز الرئيسي')", " / ".join(sorted(v)))
    if st.receipt_date:
        early = [r for r in st.rows if r.distributed_on and r.distributed_on < st.receipt_date]
        if early:
            st.issue("distributed_before_receipt", "warning", "Lines dated before the purchase order's receipt date", early[0].source_ref)


# ------------------------------------------------------------------------------------------------ parameters / pricing
def default_paper_params() -> dict:
    from pathlib import Path
    return yaml.safe_load((Path(__file__).parent / "paper_defaults.yaml").read_text(encoding="utf-8"))


def validate_paper_params(v: dict) -> dict:
    base = default_paper_params()
    unknown = [k for k in v if k not in base]
    if unknown:
        raise ValueError(f"Unknown paper parameter(s): {', '.join(unknown)}")
    for k in ("sheets_per_carton", "pages_per_sheet"):
        if k in v and (isinstance(v[k], bool) or not isinstance(v[k], (int, float)) or v[k] <= 0):
            raise ValueError(f"'{k}' must be a positive number")
    if "prices" in v:
        prev_to = None
        for p in v["prices"]:
            try:
                a, b = date.fromisoformat(p["from"]), date.fromisoformat(p["to"])
                price = float(p["price_per_carton"])
            except Exception as exc:
                raise ValueError("Each price needs from, to (YYYY-MM-DD) and price_per_carton") from exc
            if b < a or price <= 0:
                raise ValueError("A price window must end after it starts and have a positive price")
            if prev_to and a <= prev_to:
                raise ValueError("Price windows must not overlap and must be in date order")
            prev_to = b
    return v


def price_for(params: dict, on: date | None) -> dict | None:
    if on is None:
        return None
    for p in params["prices"]:
        if date.fromisoformat(p["from"]) <= on <= date.fromisoformat(p["to"]):
            return p
    return None


# ------------------------------------------------------------------------------------------------ analysis
def unit_key(r: dict) -> str:
    """A consuming unit is a branch, or one head-office department (departments are shown separately)."""
    return ("hq:" + normalize_text(r["department"] or "")) if r["is_head_office"] else r["branch_key"]


def analyze_paper(statements: list[dict], params: dict, pages_by_month: dict[tuple[int, int], dict] | None = None, filtered: bool = False) -> dict:
    """statements: [{po_no, receipt_date, rows:[{cartons, branch_display, branch_key, is_head_office, department, distributed_on}]}].
    pages_by_month: {(y,m): {"pages": int, "by_branch": {branch_key: pages}}} from the monthly consumption statements."""
    spc, pps = Decimal(str(params["sheets_per_carton"])), Decimal(str(params["pages_per_sheet"]))
    rows, notes = [], []
    for st in statements:
        pr = price_for(params, st.get("receipt_date") or next((r["distributed_on"] for r in st["rows"] if r["distributed_on"]), None))
        price = Decimal(str(pr["price_per_carton"])) if pr else None
        if price is None:
            notes.append({"code": "no_price_window", "po": st.get("po_no")})
        for r in st["rows"]:
            rows.append({**r, "po_no": st.get("po_no"), "price": price,
                         "cost": (r["cartons"] * price) if price is not None else None, "sheets": r["cartons"] * spc})
    out: dict = {"notes": notes, "n_rows": len(rows), "params": params}
    costed = [r for r in rows if r["cost"] is not None]
    out["totals"] = {"cartons": sum((r["cartons"] for r in rows), ZERO), "sheets": sum((r["sheets"] for r in rows), ZERO),
                     "cost": sum((r["cost"] for r in costed), ZERO), "cost_complete": len(costed) == len(rows),
                     "branches": len({r["branch_key"] for r in rows if not r["is_head_office"]}),
                     "hq_units": len({normalize_text(r["department"] or "") for r in rows if r["is_head_office"]})}
    months = sorted({(r["distributed_on"].year, r["distributed_on"].month) for r in rows if r["distributed_on"]})
    by_m: dict[tuple, dict] = {m: {"cartons": ZERO, "cost": ZERO, "sheets": ZERO, "branches": set()} for m in months}
    for r in rows:
        if not r["distributed_on"]:
            continue
        k = (r["distributed_on"].year, r["distributed_on"].month)
        by_m[k]["cartons"] += r["cartons"]
        by_m[k]["sheets"] += r["sheets"]
        by_m[k]["cost"] += r["cost"] or ZERO
        by_m[k]["branches"].add(r["branch_key"] if not r["is_head_office"] else "HQ:" + normalize_text(r["department"] or ""))
    pages_by_month = pages_by_month or {}
    out["by_month"] = []
    prev = None
    for m in months:
        d = by_m[m]
        row = {"period": m, "cartons": d["cartons"], "sheets": d["sheets"], "cost": d["cost"], "units": len(d["branches"]),
               "pages": None, "pages_per_carton": None, "cost_per_page": None, "need_cartons": None, "need_gap": None, "coverage": None}
        pg = None if filtered else pages_by_month.get(m)  # machine pages cover every branch; not comparable with a filtered selection
        if pg and pg.get("pages"):
            pages = Decimal(pg["pages"])
            row["pages"] = int(pages)
            row["pages_per_carton"] = float(pages / d["cartons"]) if d["cartons"] else None
            row["cost_per_page"] = (d["cost"] / pages) if out["totals"]["cost_complete"] else None
            need = pages / pps / spc
            row["need_cartons"] = need
            row["need_gap"] = d["cartons"] - need
            row["coverage"] = float(d["cartons"] / need * 100) if need else None
        if prev is not None:
            row["cartons_mom"] = d["cartons"] - prev["cartons"]
            row["cartons_mom_pct"] = float((d["cartons"] - prev["cartons"]) / prev["cartons"] * 100) if prev["cartons"] else None
        out["by_month"].append(row)
        prev = row
    # unit (branch / HQ department) x month
    units: dict[str, dict] = {}
    for r in rows:
        key = unit_key(r)
        u = units.setdefault(key, {"key": key, "name": r["branch_display"], "hq": r["is_head_office"], "department": r["department"], "cartons": ZERO, "cost": ZERO,
                                   "by_month": {}, "deliveries": 0, "flags": set()})
        u["cartons"] += r["cartons"]
        u["cost"] += r["cost"] or ZERO
        u["deliveries"] += 1
        u["flags"].update(r.get("flags", []))
        if r["distributed_on"]:
            k = (r["distributed_on"].year, r["distributed_on"].month)
            u["by_month"][k] = u["by_month"].get(k, ZERO) + r["cartons"]
    total = out["totals"]["cartons"]
    ul = sorted(units.values(), key=lambda u: -u["cartons"])
    cum = ZERO
    for i, u in enumerate(ul, 1):
        cum += u["cartons"]
        u.update({"rank": i, "share": float(u["cartons"] / total * 100) if total else None, "cum_share": float(cum / total * 100) if total else None,
                  "avg_per_month": u["cartons"] / len(months) if months else None})
    out["units"] = ul
    out["branches"] = [u for u in ul if not u["hq"]]
    out["head_office"] = sorted([u for u in ul if u["hq"]], key=lambda u: -u["cartons"])
    out["months"] = months
    hq_total = sum((u["cartons"] for u in out["head_office"]), ZERO)
    out["hq_summary"] = {"cartons": hq_total, "share": float(hq_total / total * 100) if total else None,
                         "cost": sum((u["cost"] for u in out["head_office"]), ZERO)}
    # branch-level comparison with machine pages, only for months that have a consumption statement
    out["compare"] = _compare(rows, pages_by_month, spc, pps)
    return out


def match_unit(unit: dict, machine_branches: dict[str, str]) -> tuple[str | None, str]:
    """Distribution unit -> consumption-statement branch. No guessing: both sides carry the official branch key when the branch
    register knows the name (otherwise the normalised text), and a unit matches only when the keys are identical.
    Head-office departments are not branches of the register and are never matched."""
    if unit["hq"]:
        return None, "none"
    return (unit["key"], "exact") if unit["key"] in machine_branches else (None, "none")


def _compare(rows: list[dict], pages_by_month: dict, spc: Decimal, pps: Decimal) -> list[dict]:
    out = []
    for m, pg in sorted(pages_by_month.items()):
        if not pg.get("by_branch"):
            continue
        machine_branches = {k: v for k, v in pg["by_branch"].items()}
        units: dict[str, dict] = {}
        for r in rows:
            if not r["distributed_on"] or (r["distributed_on"].year, r["distributed_on"].month) != m:
                continue
            key = unit_key(r)
            u = units.setdefault(key, {"key": key, "name": r["branch_display"], "hq": r["is_head_office"], "department": r["department"], "cartons": ZERO, "cost": ZERO})
            u["cartons"] += r["cartons"]
            u["cost"] += r["cost"] or ZERO
        for u in units.values():
            mk, how = match_unit(u, {k: k for k in machine_branches})
            pages = machine_branches.get(mk) if mk else None
            need = (Decimal(pages) / pps / spc) if pages is not None else None
            out.append({"period": m, "unit": u["name"], "hq": u["hq"], "cartons": u["cartons"], "cost": u["cost"], "match": how, "machine_branch": mk,
                        "pages": pages, "pages_per_carton": float(Decimal(pages) / u["cartons"]) if pages is not None and u["cartons"] else None,
                        "need_cartons": need, "gap": (u["cartons"] - need) if need is not None else None})
    return out
