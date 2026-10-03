"""Temporary-advance register (F4): reader, storage and analysis. Analysis only — no approval/disbursement workflow.

The register states settlement only as TEXT ("تم التسوية بتاريخ d/m/yyyy"); a settlement without a readable date is kept as
«closed, date not stated», never guessed. The monthly «اقفال» rows are the file's own control totals: the advance column is
recomputed and compared; the second hard-coded figure has no stated meaning and is shown, not used."""
import io
import re
import statistics
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal

from openpyxl import load_workbook

from app.core.cleaning.normalizers import normalize_text
from app.modules.custody_analysis.layouts import MONTHS, Issue, UnrecognisedLayout, _branch_kind

ZERO = Decimal(0)
_DATE = re.compile(r"(\d{1,2})\s*/\s*(\d{1,2})\s*/\s*(\d+)")


@dataclass
class AdvanceRow:
    source_ref: str
    holder: str | None
    branch_label: str | None
    branch_kind: str | None
    advance_on: date | None
    amount: Decimal
    purpose: str | None
    settlement_text: str | None
    settled_on: date | None
    state: str
    status_note: str | None
    flags: list[str] = field(default_factory=list)


@dataclass
class AdvanceControl:
    label: str
    month: int | None
    stated: Decimal | None
    computed: Decimal
    other_stated: Decimal | None
    source_ref: str


@dataclass
class ParsedAdvances:
    sheet: str
    rows: list[AdvanceRow] = field(default_factory=list)
    controls: list[AdvanceControl] = field(default_factory=list)
    issues: list[Issue] = field(default_factory=list)

    def issue(self, code, severity, message, example=None):
        for i in self.issues:
            if i.code == code:
                i.count += 1
                if example and len(i.examples) < 10:
                    i.examples.append(example)
                return
        self.issues.append(Issue(code, severity, message, 1, [example] if example else []))


def _dec(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    return None


def _date(v):
    return v.date() if isinstance(v, datetime) else (v if isinstance(v, date) else None)


def _settlement(text: str | None):
    """(state, settled_on, flag). text None/blank = still open."""
    if not text or not text.strip():
        return "open", None, None
    m = _DATE.search(text)
    if not m:
        return "settled_date_unreadable", None, "settlement_date_missing"
    try:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        return "settled", date(y, mo, d), None
    except ValueError:
        return "settled_date_unreadable", None, "settlement_date_invalid"


def parse_advances(wb) -> ParsedAdvances:
    ws = next((w for w in wb if any("مبلغ السلفه" in normalize_text(c.value) for row in w.iter_rows(max_row=3) for c in row if c.value)), None)
    if ws is None:
        raise UnrecognisedLayout("No advance register sheet (column 'مبلغ السلفة') found")
    hdr_row = next(r for r in range(1, 4) if any("مبلغ السلفه" in normalize_text(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)))
    cols: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        h = normalize_text(ws.cell(hdr_row, c).value)
        for key, needle in (("holder", "الاسم"), ("branch", "الفرع"), ("date", "تاريخ السلفه"), ("amount", "مبلغ السلفه"), ("purpose", "الغرض"),
                            ("settlement", "تحت التسويه"), ("status", "الموقف")):
            if needle in h and key not in cols:
                cols[key] = c
    missing = [k for k in ("branch", "date", "amount") if k not in cols]
    if missing:
        raise UnrecognisedLayout(f"Advance register columns not found: {', '.join(missing)}")
    out = ParsedAdvances(ws.title)
    g = lambda r, k: ws.cell(r, cols[k]).value if k in cols else None  # noqa: E731
    for r in range(hdr_row + 1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str) and "اقفال" in normalize_text(a):
            month = next((MONTHS[t] for t in normalize_text(a).split() if t in MONTHS), None)
            amt, other = _dec(g(r, "amount")), None
            for c in range(2, ws.max_column + 1):
                if c != cols["amount"] and _dec(ws.cell(r, c).value) is not None:
                    other = _dec(ws.cell(r, c).value)
            out.controls.append(AdvanceControl(" ".join(a.split()), month, amt, ZERO, other, f"{ws.title}!R{r}"))
            continue
        amount, dt = _dec(g(r, "amount")), _date(g(r, "date"))
        holder, branch = g(r, "holder"), g(r, "branch")
        if all(g(r, k) in (None, "") for k in cols):
            continue
        ref = f"{ws.title}!R{r}"
        if amount is None:
            out.issue("amount_missing", "critical", "A register row has no readable advance amount (row not analysed)", ref)
            continue
        flags: list[str] = []
        if dt is None:
            flags.append("advance_date_missing")
            out.issue("advance_date_missing", "warning", "An advance has no date (kept; excluded from monthly and ageing figures)", ref)
        stext = " ".join(str(g(r, "settlement")).split()) if g(r, "settlement") not in (None, "") else None
        status = " ".join(str(g(r, "status")).split()) if g(r, "status") not in (None, "") else None
        state, settled_on, flag = _settlement(stext)
        if flag:
            flags.append(flag)
            out.issue(flag, "warning", "Settlement text without a readable date (counted as closed, date not stated)", f"{ref}: {stext}")
        if state == "open" and status and "استرجاع" in normalize_text(status):
            state = "refunded_note"
            flags.append("closed_by_status_note")
            out.issue("closed_by_status_note", "info", "Closed only by a free-text note in the status column (date not stated)", f"{ref}: {status}")
        if settled_on and dt and settled_on < dt:
            flags.append("settled_before_advance")
            out.issue("settled_before_advance", "critical", "Settlement date earlier than the advance date", ref)
        label = " ".join(str(branch).split()) if branch not in (None, "") else None
        if label is None:
            out.issue("branch_missing", "warning", "An advance without a branch", ref)
        out.rows.append(AdvanceRow(ref, " ".join(str(holder).split()) if holder else None, label, _branch_kind(label) if label else None, dt, amount,
                                   " ".join(str(g(r, "purpose")).split()) if g(r, "purpose") else None, stext, settled_on, state, status, flags))
    if not out.rows:
        raise UnrecognisedLayout("The advance register has no advance rows")
    for c in out.controls:
        c.computed = sum((x.amount for x in out.rows if x.advance_on and x.advance_on.month == c.month), ZERO)
        if c.stated is not None and c.stated != c.computed:
            out.issue("control_total_differs", "warning", "A monthly closing total differs from the advances of that month", f"{c.label}: stated {c.stated} vs rows {c.computed}")
        if c.other_stated is not None:
            out.issue("closing_second_figure_unexplained", "info", "The second hard-coded figure on the monthly closing rows has no stated meaning (shown, not used)", c.label)
    seen: dict = defaultdict(list)
    for x in out.rows:
        seen[(normalize_text(x.holder), x.advance_on, x.amount)].append(x.source_ref)
    for k, refs in seen.items():
        if k[0] and len(refs) > 1:
            out.issue("possible_duplicate", "warning", "Same person, date and amount on more than one row (kept, review)", " + ".join(refs))
    return out


def as_of_of(rows: list[AdvanceRow]) -> date | None:
    ds = [d for x in rows for d in (x.advance_on, x.settled_on) if d]
    return max(ds) if ds else None


def _pct(a, b):
    return float(Decimal(a) / Decimal(b) * 100) if b else None


def _age(x, as_of):
    return (as_of - x["advance_on"]).days if as_of and x["advance_on"] else None


def analyze(rows: list[dict], controls: list[dict], th: dict, as_of: date | None, filtered: bool) -> dict:
    """rows: dicts with advance_on, amount, settled_on, state, branch_key, branch_label, branch_kind, purpose, flags, source_ref."""
    n = len(rows)
    issued = sum((x["amount"] for x in rows), ZERO)
    closed = [x for x in rows if x["state"] != "open"]
    open_ = [x for x in rows if x["state"] == "open"]
    t = {"n": n, "issued": issued, "avg": issued / n if n else None, "max": max((x["amount"] for x in rows), default=None),
         "closed_n": len(closed), "closed_amount": sum((x["amount"] for x in closed), ZERO), "open_n": len(open_), "open_amount": sum((x["amount"] for x in open_), ZERO),
         "undated_closed_n": sum(1 for x in closed if not x["settled_on"]), "no_date_n": sum(1 for x in rows if not x["advance_on"])}
    t["open_pct"] = _pct(t["open_amount"], issued)
    # ---------------------------------------------------------------- months
    ym = lambda d: (d.year, d.month)  # noqa: E731
    months = sorted({ym(x["advance_on"]) for x in rows if x["advance_on"]} | {ym(x["settled_on"]) for x in rows if x["settled_on"]})
    mrows, prev = [], None
    known = [x for x in rows if x["advance_on"] and (x["state"] == "open" or x["settled_on"])]       # rows whose open/closed timeline is fully dated
    for m in months:
        end = date(m[0] + (m[1] == 12), m[1] % 12 + 1, 1)
        iss = [x for x in rows if x["advance_on"] and ym(x["advance_on"]) == m]
        stl = [x for x in rows if x["settled_on"] and ym(x["settled_on"]) == m]
        bal = [x for x in known if x["advance_on"] < end and not (x["settled_on"] and x["settled_on"] < end)]
        c = {"period": m, "issued_n": len(iss), "issued": sum((x["amount"] for x in iss), ZERO), "settled_n": len(stl), "settled": sum((x["amount"] for x in stl), ZERO),
             "open_n": len(bal), "open": sum((x["amount"] for x in bal), ZERO)}
        c["partial"] = bool(as_of and ym(as_of) == m and end - as_of > timedelta(days=1))      # the register ends before this month does
        if prev is not None and (prev["period"][0] * 12 + prev["period"][1]) + 1 == m[0] * 12 + m[1]:
            c["d_issued"], c["d_issued_pct"] = c["issued"] - prev["issued"], _pct(c["issued"] - prev["issued"], prev["issued"])
        mrows.append(c)
        prev = c
    # ---------------------------------------------------------------- branches
    br: dict = {}
    for x in rows:
        k = x["branch_key"] or "none"
        c = br.setdefault(k, {"key": k, "label": x["branch_label"], "kind": x["branch_kind"], "n": 0, "issued": ZERO, "closed_n": 0, "closed": ZERO, "open_n": 0, "open": ZERO, "lags": [], "oldest": None})
        c["n"] += 1
        c["issued"] += x["amount"]
        if x["state"] == "open":
            c["open_n"] += 1
            c["open"] += x["amount"]
            a = _age(x, as_of)
            if a is not None:
                c["oldest"] = a if c["oldest"] is None else max(c["oldest"], a)
        else:
            c["closed_n"] += 1
            c["closed"] += x["amount"]
            if x["settled_on"] and x["advance_on"]:
                c["lags"].append((x["settled_on"] - x["advance_on"]).days)
    branches = sorted(br.values(), key=lambda c: (-c["issued"], str(c["label"])))
    for c in branches:
        c["avg_lag"] = statistics.mean(c["lags"]) if c["lags"] else None
        c["share"] = _pct(c["issued"], issued)
    # ---------------------------------------------------------------- ageing, lag, exceptions
    b1, b2, b3 = th["ageing_days_1"], th["ageing_days_2"], th["ageing_days_3"]
    bands = [("a1", 0, b1), ("a2", b1, b2), ("a3", b2, b3), ("a4", b3, None)]
    ageing = []
    for key, lo, hi in bands:
        sel = [x for x in open_ if _age(x, as_of) is not None and (_age(x, as_of) > lo or (lo == 0 and _age(x, as_of) >= 0)) and (hi is None or _age(x, as_of) <= hi)]
        ageing.append({"key": key, "lo": lo, "hi": hi, "n": len(sel), "amount": sum((x["amount"] for x in sel), ZERO)})
    lags = [(x["settled_on"] - x["advance_on"]).days for x in closed if x["settled_on"] and x["advance_on"] and "settled_before_advance" not in x["flags"]]
    lag = {"n": len(lags), "median": statistics.median(lags) if lags else None, "max": max(lags) if lags else None, "mean": statistics.mean(lags) if lags else None,
           "p90": sorted(lags)[int(round((len(lags) - 1) * 0.9))] if lags else None, "long": sum(1 for v in lags if v > th["long_settlement_days"])}
    med = statistics.median([x["amount"] for x in rows]) if rows else None
    exc: dict = {"large": [x for x in rows if med and x["amount"] > med * Decimal(str(th["large_advance_multiple"]))],
                 "long_settlement": [x for x in closed if x["settled_on"] and x["advance_on"] and (x["settled_on"] - x["advance_on"]).days > th["long_settlement_days"]],
                 "settled_before": [x for x in rows if "settled_before_advance" in x["flags"]], "closed_undated": [x for x in closed if not x["settled_on"]],
                 "no_date": [x for x in rows if not x["advance_on"]],
                 "oldest_open": sorted([x for x in open_ if _age(x, as_of) is not None and _age(x, as_of) > b3], key=lambda x: -_age(x, as_of))}
    repeating = [c for c in branches if c["n"] >= th["repeat_advances"] and c["kind"] != "group"]
    largest = sorted(rows, key=lambda x: -x["amount"])[: int(th["top_n"])]
    ctl = []
    for k in controls:
        mm = [m for m in mrows if m["period"][1] == k["month"]]
        bal = mm[0]["open"] if len(mm) == 1 else None                        # month-end open balance from the dated rows (for comparison only)
        other = Decimal(k["other_stated"]) if k.get("other_stated") is not None else None
        ctl.append({**k, "stated": Decimal(k["stated"]) if k["stated"] is not None else None, "computed": Decimal(k["computed"]), "other_stated": other,
                    "open_balance": bal, "other_vs_balance": (other - bal) if other is not None and bal is not None else None})
    controls = ctl
    return {"totals": t, "months": mrows, "branches": branches, "ageing": ageing, "lag": lag, "exceptions": exc, "repeating": repeating, "largest": largest,
            "median_advance": med, "as_of": as_of, "controls": controls, "filtered": filtered, "thresholds": th, "excluded_from_balance": sum(1 for x in rows if x not in known)}


# ------------------------------------------------------------------------------------------------ storage / build
def store(session, content: bytes, filename: str, created_by, digest: str, path: str):
    """Parse + persist one advance register as an AnalysisDataset (layout advance_register) with CustodyAdvance rows."""
    from app.core.entities import EntityResolver
    from app.models import AnalysisDataset, CustodyAdvance
    wb = load_workbook(io.BytesIO(content), data_only=True)
    p = parse_advances(wb)
    er = EntityResolver(session)
    resolved: dict[str, int | None] = {}
    unknown: set[str] = set()
    ds = AnalysisDataset(module_id="custody_analysis", layout="advance_register", scope_label="advances", file_name=filename.rsplit("/", 1)[-1], file_hash=digest,
                         storage_path=path, title=p.sheet, facts_count=len(p.rows), created_by=created_by, contains_personal_data=True,
                         period_from=min((x.advance_on for x in p.rows if x.advance_on), default=None), period_to=as_of_of(p.rows), role="advances", status="ready")
    session.add(ds)
    session.flush()
    for x in p.rows:
        key = bid = None
        if x.branch_label:
            if x.branch_kind == "head_office":
                key = "head_office"
            elif x.branch_kind == "group":
                key = "group:" + normalize_text(x.branch_label)
            else:
                n = normalize_text(x.branch_label)
                if n not in resolved:
                    resolved[n] = er.lookup("branch", x.branch_label)
                    if resolved[n] is None:
                        unknown.add(x.branch_label)
                bid = resolved[n]
                key = f"branch:{bid}" if bid else "name:" + n
        session.add(CustodyAdvance(dataset_id=ds.id, source_ref=x.source_ref, holder_text=x.holder, branch_label=x.branch_label, branch_kind=x.branch_kind, branch_key=key,
                                   branch_id=bid, advance_on=x.advance_on, amount=x.amount, purpose=x.purpose, settlement_text=x.settlement_text, settled_on=x.settled_on,
                                   state=x.state, status_note=x.status_note, flags=x.flags))
    if unknown:
        p.issue("branch_not_in_master", "info", "Branch names not found in the branch master (reported as written in the file)")
        p.issues[-1].count, p.issues[-1].examples = len(unknown), sorted(unknown)[:8]
    ds.summary = {"as_of": as_of_of(p.rows).isoformat() if as_of_of(p.rows) else None, "sheet": p.sheet,
                  "issues": [{"code": i.code, "severity": i.severity, "message": i.message, "count": i.count, "examples": i.examples} for i in p.issues],
                  "controls": [{"label": c.label, "month": c.month, "stated": str(c.stated) if c.stated is not None else None, "computed": str(c.computed),
                                "other_stated": str(c.other_stated) if c.other_stated is not None else None, "source_ref": c.source_ref} for c in p.controls]}
    session.commit()
    return ds


def load_rows(session, dataset_id: int) -> list[dict]:
    from sqlalchemy import select

    from app.models import CustodyAdvance
    return [{"source_ref": r.source_ref, "holder": r.holder_text, "branch_label": r.branch_label, "branch_kind": r.branch_kind, "branch_key": r.branch_key,
             "advance_on": r.advance_on, "amount": Decimal(r.amount), "purpose": r.purpose, "settlement_text": r.settlement_text, "settled_on": r.settled_on,
             "state": r.state, "flags": list(r.flags or [])}
            for r in session.scalars(select(CustodyAdvance).where(CustodyAdvance.dataset_id == dataset_id).order_by(CustodyAdvance.id))]


def delete(session, dataset_id: int) -> None:
    from sqlalchemy import delete as sql_delete

    from app.models import AnalysisDataset, CustodyAdvance
    session.execute(sql_delete(CustodyAdvance).where(CustodyAdvance.dataset_id == dataset_id))
    session.execute(sql_delete(AnalysisDataset).where(AnalysisDataset.id == dataset_id))
    session.commit()


def build(session, ds, lang: str, admin: bool, filters: dict):
    from app.core.analysis import AnalysisError
    from app.core.settings_store import effective_thresholds
    from app.modules.custody_analysis import advances_report
    th, origin = effective_thresholds(session, "custody_advances")
    rows = load_rows(session, ds.id)
    as_of = as_of_of_dicts(rows)
    ym = lambda d: d.strftime("%Y-%m") if d else None  # noqa: E731
    sel = rows
    f = {k: v for k, v in (filters or {}).items() if v}
    if f.get("periods"):
        sel = [x for x in sel if ym(x["advance_on"]) in set(f["periods"])]
    if f.get("branches"):
        sel = [x for x in sel if x["branch_key"] in set(f["branches"])]
    if f.get("states"):
        sel = [x for x in sel if ("open" if x["state"] == "open" else "closed") in set(f["states"])]
    if f and not sel:
        raise AnalysisError(409, "The selected filters match no advances")
    controls = (ds.summary or {}).get("controls", [])
    a = analyze(sel, controls, th, as_of, bool(f))
    a["thresholds_origin"] = origin
    rm = advances_report.build_report(ds, a, rows, lang, admin, th, origin, f, (ds.summary or {}).get("issues", []))
    return rm, a


def as_of_of_dicts(rows: list[dict]) -> date | None:
    ds = [d for x in rows for d in (x["advance_on"], x["settled_on"]) if d]
    return max(ds) if ds else None
