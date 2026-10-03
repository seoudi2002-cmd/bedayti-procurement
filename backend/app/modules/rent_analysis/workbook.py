"""Reader for the rent-contract workbook: one sheet per governorate (+ Head Office), contracts in rows, monthly rent in columns.

What the file shows (and how it is read, never fixed):
* The governorate sheets repeat one master list of contracts; each sheet's own footer total adds up *its* contracts with a
  formula (=SUM(AO4+AO5+...)), which is the file's explicit statement of which contracts belong to the governorate. Copies of a
  contract in other sheets are stale duplicates and are not added to the totals.
* Blank «_____» means no payment recorded that month (not zero). Columns are located by header text, not by letter.
* Dates are text d/m/yyyy; unreadable ones are reported and kept as written, never repaired."""
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.core.analysis.opsupport import Issues, fold, month_in_text, num, period_str

_DATE = re.compile(r"^\s*(\d{1,2})\s*[/\-.]\s*(\d{1,2})\s*[/\-.]\s*(\d{4})\s*$")
_REF = re.compile(r"[A-Z]{1,3}(\d+)")
MARKER = re.compile(r"^[\s_\-–—.]+$")


class NotARentWorkbook(Exception):
    pass


@dataclass
class Copy:
    sheet: str
    row: int
    name: str
    hq: bool
    footer: bool = False
    start: date | None = None
    start_raw: str | None = None
    end: date | None = None
    end_raw: str | None = None
    advance: float | None = None
    deposit: float | None = None
    contract_rent: float | None = None
    current_rent: float | None = None
    landlord: str | None = None
    months: dict = field(default_factory=dict)      # period -> float
    nopay: set = field(default_factory=set)         # periods written as «_____»
    flags: list = field(default_factory=list)

    @property
    def key(self) -> str:
        return f"{fold(self.name)}|{self.start.isoformat() if self.start else (self.start_raw or '?')}"


@dataclass
class SheetControl:
    sheet: str
    row: int
    label: str | None
    values: dict                                    # period -> stated number
    owner_rows: set                                 # rows its formula adds up (empty for plain totals)


@dataclass
class Parsed:
    copies: list = field(default_factory=list)
    controls: list = field(default_factory=list)
    sheets: list = field(default_factory=list)      # profile rows
    hidden: list = field(default_factory=list)
    issues: Issues = field(default_factory=Issues)


def _clean(s) -> str:
    return " ".join(str(s).replace("\n", " ").split())


def _date(v, flags: list, what: str):
    if isinstance(v, datetime):
        flags.append(f"{what}_date_cell")
        return v.date(), None
    if isinstance(v, date):
        return v, None
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, None
    raw = str(v).strip()
    m = _DATE.match(raw)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            return date(y, mo, d), raw
        except ValueError:
            pass
    flags.append(f"{what}_unreadable")
    return None, raw


def is_rent_workbook(content: bytes) -> bool:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True)
    except Exception:  # noqa: BLE001
        return False
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(min_row=1, max_row=5, values_only=True):
                t = " ".join(fold(c) for c in row if c is not None)
                if "بدايه العقد" in t and "نهايه العقد" in t:
                    return True
    finally:
        wb.close()
    return False


def _header(ws):
    for r in range(1, 6):
        cells = {c.column: fold(c.value) for c in ws[r] if c.value is not None}
        if any("بدايه العقد" in t for t in cells.values()):
            return r, cells
    return None, {}


def _month_columns(ws, hr: int, issues: Issues, sheet: str, seen: set):
    """[(col, period)] following the header sequence: a header with no year continues the previous one; a stated year that goes backwards
    (a typo) is placed by sequence and reported."""
    out, prev = [], None
    for c in ws[hr]:
        if c.value is None:
            continue
        m, y = month_in_text(c.value)
        if m is None or fold(c.value).startswith("القيمه"):
            continue
        if y is None:
            if prev is None:
                issues.add("month_header_unplaced", "warning", "A month header with no year and nothing before it (column ignored)", f"{sheet}!{c.coordinate}")
                continue
            y = prev[0] if m > prev[1] else prev[0] + 1
            if ("noyear", _clean(c.value)) not in seen:
                seen.add(("noyear", _clean(c.value)))
                issues.add("month_header_without_year", "info", "Month headers written without a year are placed after the previous month's year", f"{sheet}!{c.coordinate}: {_clean(c.value)} → {y}")
        elif prev and (y, m) <= prev:
            exp = (prev[0] + 1, 1) if prev[1] == 12 else (prev[0], prev[1] + 1)
            if m == exp[1]:
                if ("yr", sheet, c.coordinate) not in seen:
                    seen.add(("yr", sheet, c.coordinate))
                    issues.add("month_header_year_inconsistent", "warning", "A month header whose year goes backwards; placed by sequence (not corrected in the file)", f"{sheet}!{c.coordinate}: «{_clean(c.value)}» → {period_str(*exp)}")
                y = exp[0]
            else:
                issues.add("month_header_out_of_sequence", "warning", "A month header out of sequence (used as written)", f"{sheet}!{c.coordinate}: {_clean(c.value)}")
        prev = (y, m)
        out.append((c.column, period_str(y, m)))
    return out


def parse(content: bytes) -> Parsed:
    wbv = load_workbook(io.BytesIO(content), data_only=True)
    wbf = load_workbook(io.BytesIO(content))
    p = Parsed()
    any_sheet = False
    seen: set = set()
    for ws in wbv.worksheets:
        if ws.sheet_state != "visible":
            p.hidden.append(ws.title)
            continue
        hr, hdr = _header(ws)
        if hr is None:
            p.issues.add("sheet_unrecognised", "info", "Visible sheet without a contract header row (not read)", ws.title)
            p.sheets.append({"sheet": ws.title, "read": False, "contracts": 0})
            continue
        any_sheet = True
        wf = wbf[ws.title]
        cols: dict[str, int] = {}
        for col, t in hdr.items():
            k = None
            if "المركز" in t:
                k = "name"
            elif "المقدم" in t:
                k = "advance"
            elif "التامين" in t:
                k = "deposit"
            elif "المالك" in t:
                k = "landlord"
            elif "بدايه العقد" in t:
                k = "start"
            elif "نهايه العقد" in t:
                k = "end"
            elif "القيمه الايجاريه" in t:
                k = "current" if ("الحاليه" in t or "الفعليه" in t) else "contract"
            if k is None:
                continue
            if k in cols:     # a structural header repeated further right (typically inside the month region): not a second field
                if cols[k] != col:
                    p.issues.add("repeated_header", "warning", "A structural header repeated inside the sheet (the first one is used; the repeat's values are not assigned to a month)", f"{ws.title}!{get_column_letter(col)}: {_clean(ws.cell(hr, col).value)}")
                continue
            cols[k] = col
        if "name" not in cols or "start" not in cols:
            p.issues.add("sheet_unrecognised", "info", "Visible sheet without a contract name/start column (not read)", ws.title)
            p.sheets.append({"sheet": ws.title, "read": False, "contracts": 0})
            continue
        months = _month_columns(ws, hr, p.issues, ws.title, seen)
        mcols = {c for c, _p in months}
        first_m, last_m = (min(mcols), max(mcols)) if mcols else (0, 0)
        # month-region columns without any header that nevertheless hold values
        used = set(cols.values())
        unlabeled = [c for c in range(first_m, last_m + 1) if c not in mcols and c not in used]
        hq = "landlord" in cols
        first_control = None
        n_contracts = 0
        unl: dict[int, int] = {}
        for r in range(hr + 1, ws.max_row + 1):
            raw_name = ws.cell(r, cols["name"]).value
            name = _clean(raw_name) if raw_name is not None else ""
            vals = {col: ws.cell(r, col).value for col, _p in months}
            has_vals = any(v is not None for v in vals.values())
            formulas = [wf.cell(r, c).value for c in range(first_m, last_m + 1) if isinstance(wf.cell(r, c).value, str) and wf.cell(r, c).value.startswith("=")] if mcols else []
            if not name or fold(name).startswith(("اجمالي", "الاجمالي")):
                if has_vals or formulas:
                    owner_rows = set()
                    for f in formulas:
                        if f.upper().startswith("=SUM(") and "+" in f:
                            owner_rows |= {int(x) for x in _REF.findall(f) if int(x) < r}
                    if not name and first_control is None:
                        first_control = r
                    p.controls.append(SheetControl(ws.title, r, name or None, {pe: num(vals[c]) for c, pe in months if num(vals[c]) is not None}, owner_rows))
                continue
            flags: list[str] = []
            cp = Copy(sheet=ws.title, row=r, name=name, hq=hq, footer=first_control is not None and r > first_control)
            cp.start, cp.start_raw = _date(ws.cell(r, cols["start"]).value, flags, "start")
            if "end" in cols:
                cp.end, cp.end_raw = _date(ws.cell(r, cols["end"]).value, flags, "end")
            for f, k in (("advance", "advance"), ("deposit", "deposit"), ("contract_rent", "contract"), ("current_rent", "current")):
                if k in cols:
                    setattr(cp, f, num(ws.cell(r, cols[k]).value))
            if "landlord" in cols and ws.cell(r, cols["landlord"]).value:
                cp.landlord = _clean(ws.cell(r, cols["landlord"]).value)
            for col, pe in months:
                v = vals[col]
                n = num(v)
                if n is not None:
                    cp.months[pe] = n
                elif isinstance(v, str) and v.strip() and MARKER.match(v):
                    cp.nopay.add(pe)
                elif v is not None and str(v).strip():
                    cp.flags.append(f"non_numeric:{pe}:{_clean(v)[:20]}")
            for c in unlabeled:
                if num(ws.cell(r, c).value) is not None:
                    unl[c] = unl.get(c, 0) + 1
            cp.flags = flags + cp.flags
            p.copies.append(cp)
            n_contracts += 1
        for c, n in unl.items():
            p.issues.add("unlabelled_value_column", "warning", "A month-region column without a header that holds values (not assigned to any month; its values are not used)", f"{ws.title}!{get_column_letter(c)} ({n} rows)")
        p.sheets.append({"sheet": ws.title, "read": True, "contracts": n_contracts, "header_row": hr, "hq": hq,
                         "months": [months[0][1], months[-1][1]] if months else None})
    if not any_sheet:
        raise NotARentWorkbook()
    return p


def bracket(name: str) -> str | None:
    m = re.search(r"[\(（]\s*([^()）]*?)\s*[\)）]", name)
    return m.group(1).strip() if m and m.group(1).strip() else None


def assemble(p: Parsed) -> dict:
    """Pick one authoritative copy per contract (see module docstring) and resolve its governorate from explicit evidence only."""
    by_key: dict[str, list[Copy]] = {}
    for c in p.copies:
        by_key.setdefault(c.key, []).append(c)
    owners: dict[str, set] = {}
    for ct in p.controls:
        if ct.owner_rows:
            owners.setdefault(ct.sheet, set()).update(ct.owner_rows)
    sheet_names = [s["sheet"] for s in p.sheets if s["read"] and not s.get("hq")]
    contracts, excluded = [], []
    stale_cells = stale_contracts = 0
    for key, copies in by_key.items():
        if len(copies) == 1 and copies[0].hq:
            ch = copies[0]
            gov, ev = ch.sheet, "sheet"
            copies_other = []
        else:
            owner = [c for c in copies if c.row in owners.get(c.sheet, set())]
            gov = ev = None
            ch = None
            if owner:
                ch, gov, ev = owner[0], owner[0].sheet, "sheet_formula"
                if len({c.sheet for c in owner}) > 1:
                    p.issues.add("multiple_owner_sheets", "warning", "A contract that more than one sheet's footer total adds up (first used)", ch.name)
            br = bracket(copies[0].name)
            if ch is None and br:
                match = next((s for s in sheet_names if fold(s) == fold(br)), None)
                cand = next((c for c in copies if c.sheet == match), None) if match else None
                if cand:
                    ch, gov, ev = cand, match, "name_text"
            if ch is None:
                foot = [c for c in copies if c.footer]
                if len(copies) == 1 and foot:
                    ch, gov, ev = copies[0], copies[0].sheet, "sheet_footer"
            if ch is None:
                ne = [c for c in copies if c.months or c.nopay]
                same_all = all((c.months, c.nopay) == (ne[0].months, ne[0].nopay) for c in ne) if ne else True
                if same_all:
                    ch = (ne or copies)[0]
                    gov = None
                    ev = None
                else:
                    excluded.append({"name": copies[0].name, "reason": "copies_conflict_unowned", "sheets": sorted({c.sheet for c in copies})})
                    p.issues.add("copies_conflict_unowned", "critical", "A contract with differing copies in several sheets and no sheet that adds it up: not valued, needs review", copies[0].name)
                    continue
            if gov and br and fold(br) != fold(gov) and ev == "sheet_formula":
                p.issues.add("governorate_conflict", "warning", "The governorate written in the contract name differs from the sheet whose total adds it up (the sheet is used)", f"{ch.name} ≠ {gov}")
            copies_other = [c for c in copies if c is not ch]
        diff = 0
        for o in copies_other:
            diff += sum(1 for pe, v in o.months.items() if pe in ch.months and abs(v - ch.months[pe]) > 0.005)
            diff += sum(1 for pe in o.nopay if pe in ch.months)
            diff += sum(1 for pe, v in ch.months.items() if pe in o.nopay)
        if diff:
            stale_cells += diff
            stale_contracts += 1
        for fl in ch.flags:
            if fl.endswith("_unreadable"):
                p.issues.add(fl, "warning", f"A {fl.split('_')[0]} date that is not a valid date (kept as written, not repaired)", f"{ch.name}: {ch.start_raw if fl.startswith('start') else ch.end_raw}")
            elif fl.endswith("_date_cell"):
                p.issues.add("date_cell", "info", "Dates stored as real date cells: the day/month order cannot be confirmed from the cell (shown as stored)", ch.name)
            elif fl.startswith("non_numeric:"):
                p.issues.add("non_numeric_value", "warning", "A monthly cell that is neither a number nor the «no payment» marker (ignored)", f"{ch.name}: {fl[12:]}")
        if gov is None:
            p.issues.add("governorate_unresolved", "warning", "Contracts whose governorate has no explicit evidence (neither a sheet total that adds them up nor a governorate in the name): shown Unallocated", ch.name)
        contracts.append({"key": key, "copy": ch, "governorate": gov, "evidence": ev, "copies": len(copies)})
    if stale_cells:
        p.issues.add("stale_copies", "info", "Contracts whose copies in the other governorate sheets differ from the sheet that adds them up (those copies are not used)", f"{stale_contracts} contracts, {stale_cells} cells")
        p.issues.items[-1].count = stale_contracts
    return {"contracts": contracts, "excluded": excluded, "stale_cells": stale_cells, "stale_contracts": stale_contracts}
