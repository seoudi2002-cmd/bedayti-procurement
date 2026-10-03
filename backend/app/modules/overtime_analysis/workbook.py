"""Reader for the monthly overtime statement: one sheet per month, one row per employee (code + name), columns for missions (day/night),
meals, overtime hours (day / night, each with its weighted value) and totals; plus an annual `report` sheet that adds the months with formulas.

Read as stated, never fixed:
* The period of a sheet is the month and year written in its TITLE (sheet names are only labels: here two are swapped and one repeats another).
* Blank employee rows mean no entry for the month (never zero). Hours are the statement's own quantities; there are no amounts, no salary and
  no attendance data in the file, so nothing is converted to money.
* Weighted values are checked against the multipliers the file itself states (header cells 1.35 / 1.7); differences are reported."""
import io
import re
from collections import Counter
from dataclasses import dataclass, field

from openpyxl import load_workbook

from app.core.analysis.opsupport import Issues, fold, month_in_text, num, period_str

FIELDS = ("mission_day", "mission_night", "meals", "day_hours", "day_weighted", "night_hours", "night_weighted", "raw_total", "weighted_total", "equal_pay_days", "equal_pay_x2")


class NotAnOvertimeWorkbook(Exception):
    pass


@dataclass
class Row:
    code: str
    name: str
    values: dict
    ref: str


@dataclass
class SheetData:
    sheet: str
    period: str | None
    title: str
    rows: list = field(default_factory=list)        # Row (every listed employee, entry or not)
    factors: tuple | None = None                    # (day, night) multipliers stated in the sheet header
    totals_cols: int = 1


@dataclass
class Parsed:
    sheets: list = field(default_factory=list)
    summary_controls: dict = field(default_factory=dict)   # field -> stated annual total
    factors: tuple | None = None
    excluded: list = field(default_factory=list)           # sheets left out of the analysis: {sheet, period, reason, same_as}
    issues: Issues = field(default_factory=Issues)


def _find(row_cells: dict, *tokens):
    for col, t in row_cells.items():
        if all(x in t for x in tokens):
            return col
    return None


def is_overtime_workbook(content: bytes) -> bool:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True)
    except Exception:  # noqa: BLE001
        return False
    try:
        for ws in wb.worksheets:
            t = " ".join(fold(c) for row in ws.iter_rows(min_row=1, max_row=4, values_only=True) for c in row if c is not None)
            if "الكود" in t and "الساعات الاضافيه" in t:
                return True
    finally:
        wb.close()
    return False


def parse(content: bytes, year: int | None = None) -> Parsed:
    wbv = load_workbook(io.BytesIO(content), data_only=True)
    wbf = load_workbook(io.BytesIO(content))
    p = Parsed()
    summary = None
    for ws in wbv.worksheets:
        hdr_row = next((r for r in range(1, 6) if any(fold(c.value) == "الكود" for c in ws[r])), None)
        if hdr_row is None:
            continue
        wf = wbf[ws.title]
        first_data = hdr_row + 2
        is_summary = isinstance(wf.cell(first_data, 3).value, str) and str(wf.cell(first_data, 3).value).startswith("=") and "!" in str(wf.cell(first_data, 3).value)
        cells = {c.column: fold(c.value) for c in ws[hdr_row] if c.value is not None}
        sub = {c.column: c.value for c in ws[hdr_row + 1] if c.value is not None}
        c_code, c_name = _find(cells, "الكود"), _find(cells, "الاسم")
        c_mis, c_meal, c_hrs, c_eq = _find(cells, "مأموريات") or _find(cells, "ماموريات"), _find(cells, "وجبات"), _find(cells, "الساعات"), _find(cells, "مثل")
        if None in (c_code, c_name, c_mis, c_meal, c_hrs, c_eq):
            p.issues.add("sheet_unrecognised", "info", "A sheet with the overtime headers but a layout the reader does not know (not read)", ws.title)
            continue
        n_tot = c_eq - (c_hrs + 4)
        cols = {"mission_day": c_mis, "mission_night": c_mis + 1, "meals": c_meal, "day_hours": c_hrs, "day_weighted": c_hrs + 1, "night_hours": c_hrs + 2, "night_weighted": c_hrs + 3,
                "equal_pay_days": c_eq, "equal_pay_x2": c_eq + 1}
        if n_tot >= 2:
            cols["raw_total"], cols["weighted_total"] = c_hrs + 4, c_hrs + 5
        elif n_tot == 1:
            cols["weighted_total"] = c_hrs + 4
        fac = (num(sub.get(c_hrs + 1)), num(sub.get(c_hrs + 3)))
        factors = fac if all(f for f in fac) else None
        if is_summary:
            summary = ws
            if factors:
                p.factors = factors
            for r in range(first_data, ws.max_row + 1):
                a = ws.cell(r, c_code).value
                if isinstance(a, str) and "اجمالي" in fold(a):
                    p.summary_controls = {f: num(ws.cell(r, col).value) for f, col in cols.items() if num(ws.cell(r, col).value) is not None}
            continue
        title = ""
        for r in range(1, hdr_row):
            for c in ws[r]:
                if isinstance(c.value, str) and c.value.strip():
                    title = " ".join(c.value.split())
                    break
            if title:
                break
        m, y = month_in_text(title)
        sd = SheetData(sheet=ws.title, period=None, title=title, factors=factors, totals_cols=n_tot)
        if m is None:
            p.issues.add("period_unreadable", "warning", "A monthly sheet whose title states no month (not read)", ws.title)
            continue
        if y is None:
            y = year
        if y is None:
            p.issues.add("year_missing", "warning", "A monthly sheet whose title states no year and none was supplied (not read)", ws.title)
            continue
        sd.period = period_str(y, m)
        mm_name, _y = month_in_text(ws.title)
        if mm_name is not None and mm_name != m:
            p.issues.add("sheet_name_title_mismatch", "warning", "The sheet's name and the month written in its title differ (the title is used)", f"{ws.title} → {period_str(y, m)}")
        for r in range(first_data, ws.max_row + 1):
            code = ws.cell(r, c_code).value
            name = ws.cell(r, c_name).value
            if code is None and name is None:
                continue
            if isinstance(code, str) and "اجمالي" in fold(code):
                continue
            vals = {f: num(ws.cell(r, col).value) for f, col in cols.items()}
            sd.rows.append(Row(code=str(int(code)) if isinstance(code, (int, float)) else str(code).strip(), name=" ".join(str(name or "").split()), values=vals, ref=f"{ws.title}!{r}"))
        p.sheets.append(sd)
    if not p.sheets:
        raise NotAnOvertimeWorkbook()
    if summary is None:
        p.issues.add("no_summary_sheet", "info", "No annual summary sheet (nothing to reconcile the months against)", "")
    return p


def _sig(sd: SheetData):
    return tuple((r.code, tuple(sorted((k, v) for k, v in r.values.items() if v is not None))) for r in sd.rows)


def select(p: Parsed) -> list[SheetData]:
    """One sheet per period: a sheet that repeats a period already read is excluded (identical = a copy; different = a conflict)."""
    first: dict[str, SheetData] = {}
    out = []
    for sd in p.sheets:
        if sd.period in first:
            kept = first[sd.period]
            if _sig(sd) == _sig(kept):
                p.excluded.append({"sheet": sd.sheet, "period": sd.period, "reason": "duplicate_period_sheet", "same_as": kept.sheet, "employees": len(sd.rows)})
                p.issues.add("duplicate_period_sheet", "warning", "A sheet repeats, row for row, the data of another sheet under the same month in its title (excluded; the month it was named for has no data of its own)", f"{sd.sheet} = {kept.sheet} ({sd.period})")
            else:
                p.excluded.append({"sheet": sd.sheet, "period": sd.period, "reason": "conflicting_period_sheets", "same_as": kept.sheet, "employees": len(sd.rows)})
                p.issues.add("conflicting_period_sheets", "critical", "Two sheets state the same month with different data (the first is used; needs review)", f"{kept.sheet} / {sd.sheet} ({sd.period})")
            continue
        first[sd.period] = sd
        out.append(sd)
    # multipliers stated by the file
    stated = Counter(sd.factors for sd in p.sheets if sd.factors)
    if p.factors is None and stated:
        p.factors = stated.most_common(1)[0][0]
    return out


def check(p: Parsed, sheets: list[SheetData]) -> None:
    f = p.factors
    for sd in sheets:
        for r in sd.rows:
            v = r.values
            if f and v["day_hours"] is not None and v["day_weighted"] is not None and abs(v["day_weighted"] - v["day_hours"] * f[0]) > 0.01:
                p.issues.add("weighted_differs_from_multiplier", "warning", "A weighted value that differs from hours × the multiplier the file states", f"{r.ref} (day)")
            if f and v["night_hours"] is not None and v["night_weighted"] is not None and abs(v["night_weighted"] - v["night_hours"] * f[1]) > 0.01:
                p.issues.add("weighted_differs_from_multiplier", "warning", "A weighted value that differs from hours × the multiplier the file states", f"{r.ref} (night)")
            if v["weighted_total"] is not None and v["day_weighted"] is not None and v["night_weighted"] is not None and abs(v["weighted_total"] - v["day_weighted"] - v["night_weighted"]) > 0.01:
                p.issues.add("total_differs_from_parts", "warning", "A stated total that differs from the sum of its parts", r.ref)
