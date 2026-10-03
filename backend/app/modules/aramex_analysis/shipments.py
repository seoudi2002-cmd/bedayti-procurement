"""Aramex shipment detail workbook (one invoice). Rows with a numeric AWB are shipments; a row whose AWB is text is an
invoice-level adjustment (e.g. 'TAX Rounding Diff'), never a shipment; the unlabeled numeric row is the sheet's own total.

Column meaning, as observed and stated in the report: 'Net Value' is the line total INCLUDING tax; 'Amount' repeats the base
charge; 'Airwaybill Date' is not the pick-up date. None of this is corrected: the PDF supplies the dates."""
import io
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

ZERO = Decimal(0)


def normalize_text(v) -> str:   # header names: case and spacing only
    return " ".join(str(v).lower().split()) if v is not None else ""


class UnrecognisedShipments(ValueError):
    pass


@dataclass
class XlsxShipment:
    source_ref: str
    awb: str
    origin: str | None
    destination: str | None
    shipper_name: str | None
    sent_by: str | None
    consignee_name: str | None
    attention: str | None
    base: Decimal
    other: Decimal
    tax: Decimal
    gross: Decimal
    weight: Decimal | None
    actual_weight: Decimal | None


@dataclass
class Issue:
    code: str
    severity: str
    message: str
    count: int = 1
    examples: list[str] = field(default_factory=list)


@dataclass
class AramexXlsx:
    bill_doc: str | None = None
    shipments: list[XlsxShipment] = field(default_factory=list)
    adjustments: list[dict] = field(default_factory=list)       # {label, tax, net, source_ref}
    stated_totals: dict[str, Decimal] = field(default_factory=dict)
    awb_date: date | None = None
    issues: list[Issue] = field(default_factory=list)

    def issue(self, code, severity, message, example=None):
        for i in self.issues:
            if i.code == code:
                i.count += 1
                if example and len(i.examples) < 10:
                    i.examples.append(example)
                return
        self.issues.append(Issue(code, severity, message, 1, [example] if example else []))


_COLS = {"bill_doc": "bill. doc.", "awb": "airway bill no.", "origin": "origin location", "destination": "destination location",
         "shipper_name": "shipper name", "sent_by": "shipper sent by", "consignee_name": "consignee name", "attention": "consignee attention by",
         "base": "base charge", "other": "other charge", "gross": "net value", "tax": "tax amount", "amount": "amount",
         "weight": "chargeable weight", "actual": "actual weight", "awb_date": "airwaybill date"}


def _dec(v) -> Decimal | None:
    if v is None or v == "" or isinstance(v, bool):
        return None
    try:
        return Decimal(str(v))
    except Exception:
        return None


def _txt(v) -> str | None:
    t = " ".join(str(v).split()) if v is not None else ""
    return t or None


def is_shipments_workbook(content: bytes) -> bool:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return any("airway bill no." in {normalize_text(c) for c in row if c is not None} for row in ws.iter_rows(max_row=3, values_only=True))
    except Exception:
        return False


def parse_shipments_xlsx(content: bytes) -> AramexXlsx:
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.worksheets[0]
    col: dict[str, int] = {}
    hrow = None
    for r in range(1, 6):
        names = {normalize_text(ws.cell(r, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(r, c).value is not None}
        if "airway bill no." in names:
            col, hrow = {k: names[v] for k, v in _COLS.items() if v in names}, r
            break
    missing = [k for k in ("awb", "base", "other", "tax", "gross") if k not in col]
    if hrow is None or missing:
        raise UnrecognisedShipments("Not an Aramex shipment detail sheet (Airway Bill No. / charges columns not found)")
    out = AramexXlsx()
    docs: set[str] = set()
    dates: set[date] = set()
    all_tax_incl = all_amount_base = True
    for r in range(hrow + 1, ws.max_row + 1):
        def g(k, r=r):
            return ws.cell(r, col[k]).value if k in col else None
        awb = _txt(g("awb"))
        nums = {k: _dec(g(k)) for k in ("base", "other", "tax", "gross")}
        bill = _txt(g("bill_doc"))
        if awb is None:
            if any(v is not None for v in nums.values()):          # the sheet's own total row
                out.stated_totals = {k: v or ZERO for k, v in nums.items()}
            continue
        if bill:
            docs.add(bill)
        ref = f"{ws.title}!R{r}"
        if not awb.isdigit():                                      # adjustment line, not a shipment
            out.adjustments.append({"label": awb, "tax": nums["tax"] or ZERO, "net": nums["gross"] or ZERO, "base": nums["base"] or ZERO,
                                    "other": nums["other"] or ZERO, "source_ref": ref})
            continue
        if any(nums[k] is None for k in nums):
            out.issue("amount_missing", "critical", "A shipment row has a missing amount", ref)
        n = {k: v or ZERO for k, v in nums.items()}
        if abs(n["base"] + n["other"] + n["tax"] - n["gross"]) > Decimal("0.02"):
            all_tax_incl = False
        if _dec(g("amount")) != n["base"]:
            all_amount_base = False
        if d := g("awb_date"):
            dates.add(d.date() if hasattr(d, "date") else d)
        out.shipments.append(XlsxShipment(ref, awb, _txt(g("origin")), _txt(g("destination")), _txt(g("shipper_name")), _txt(g("sent_by")),
                                          _txt(g("consignee_name")), _txt(g("attention")), n["base"], n["other"], n["tax"], n["gross"],
                                          _dec(g("weight")), _dec(g("actual"))))
    if not out.shipments:
        raise UnrecognisedShipments("No shipment rows found")
    if len(docs) > 1:
        raise UnrecognisedShipments("The sheet holds more than one invoice (Bill. Doc.); upload one invoice per file")
    out.bill_doc = next(iter(docs), None)
    if len(dates) == 1:
        out.awb_date = next(iter(dates))
    seen: dict[str, int] = {}
    for s in out.shipments:
        seen[s.awb] = seen.get(s.awb, 0) + 1
    for awb, c in seen.items():
        if c > 1:
            out.issue("duplicate_awb", "critical", "The same AWB appears on more than one row", awb)
    if all_tax_incl:
        out.issue("net_value_includes_tax", "info", "'Net Value' equals base + other + tax on every row: it is the total INCLUDING tax (the invoice's 'Net Amount' excludes tax)")
    if all_amount_base:
        out.issue("amount_equals_base", "info", "'Amount' repeats the base charge on every row (not used)")
    if out.awb_date:
        out.issue("awb_date_not_pickup", "info", "'Airwaybill Date' holds one value for every row; the pick-up date from the invoice PDF is used instead", out.awb_date.isoformat())
    if out.adjustments:
        out.issue("adjustment_rows", "info", "Invoice-level adjustment rows are kept apart from the shipments", "; ".join(a["label"] for a in out.adjustments))
    if out.stated_totals:
        for k in ("base", "other", "tax", "gross"):
            tot = sum((getattr(s, k) for s in out.shipments), ZERO) + sum((a.get(k if k != "gross" else "net", ZERO) for a in out.adjustments), ZERO)
            if tot != out.stated_totals.get(k, ZERO):
                out.issue("total_row_differs", "critical", f"The sheet's total row differs from its rows ({k})", f"rows {tot} vs stated {out.stated_totals.get(k)}")
    return out

