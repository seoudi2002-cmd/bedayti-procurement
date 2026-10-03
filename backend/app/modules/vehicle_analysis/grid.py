"""A uniform read-only view of a worksheet from .xls (xlrd) or .xlsx (openpyxl): cells by 0-based (row, col), values as written
(empty = None), spreadsheet dates converted, and — for .xlsx only — the formula text of a cell."""
import io
from datetime import date, datetime, timedelta

from openpyxl import load_workbook


def excel_date(v, datemode: int = 0) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:
        base = date(1904, 1, 1) if datemode else date(1899, 12, 30)
        return base + timedelta(days=int(v))
    return None


class Grid:
    def __init__(self, title, nrows, ncols, getter, formula=None, merged=None, hidden=False, datemode=0):
        self.title, self.nrows, self.ncols, self._get, self._formula, self.merged, self.hidden, self.datemode = title, nrows, ncols, getter, formula, merged or [], hidden, datemode

    def v(self, r: int, c: int):
        if r < 0 or c < 0 or r >= self.nrows or c >= self.ncols:
            return None
        x = self._get(r, c)
        return None if x == "" else x

    def formula(self, r: int, c: int):
        return self._formula(r, c) if self._formula else None

    def text(self, r: int, c: int) -> str:
        x = self.v(r, c)
        return " ".join(str(x).split()) if x is not None else ""

    def row_texts(self, r: int) -> dict:
        return {c: self.text(r, c) for c in range(self.ncols) if self.v(r, c) is not None}

    def group_fill(self, r: int, c: int):
        """The text of cell (r, c), or of the merged range that covers it (merged cells hold their value in the top-left cell)."""
        x = self.v(r, c)
        if x is not None:
            return " ".join(str(x).split())
        for (r0, c0, r1, c1) in self.merged:
            if r0 <= r < r1 and c0 <= c < c1:
                return self.text(r0, c0)
        return ""


def grids(content: bytes, filename: str) -> list[Grid]:
    name = filename.lower()
    out = []
    if name.endswith(".xls"):
        import xlrd
        wb = xlrd.open_workbook(file_contents=content, formatting_info=False)
        for s in wb.sheets():
            out.append(Grid(s.name, s.nrows, s.ncols, (lambda s: lambda r, c: s.cell_value(r, c))(s), None, [],
                            s.visibility != 0, wb.datemode))
        return out
    wbv = load_workbook(io.BytesIO(content), data_only=True)
    wbf = load_workbook(io.BytesIO(content))
    for ws in wbv.worksheets:
        wf = wbf[ws.title]
        merged = [(m.min_row - 1, m.min_col - 1, m.max_row, m.max_col) for m in ws.merged_cells.ranges]
        f = (lambda wf: lambda r, c: (wf.cell(r + 1, c + 1).value if isinstance(wf.cell(r + 1, c + 1).value, str) and str(wf.cell(r + 1, c + 1).value).startswith("=") else None))(wf)
        out.append(Grid(ws.title, ws.max_row, ws.max_column, (lambda ws: lambda r, c: ws.cell(r + 1, c + 1).value)(ws), f, merged, ws.sheet_state != "visible"))
    return out
