"""Read CSV/XLSX uploads into a uniform table, tolerating the usual Excel mess
(title rows above the header, blank rows, Arabic CSV encodings)."""
import csv
import io
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import PurePath

from openpyxl import load_workbook

from app.core.cleaning.normalizers import clean_text

HEADER_SCAN_ROWS = 20


class UnsupportedFileError(ValueError):
    pass


@dataclass
class TableData:
    headers: list[str]
    rows: list[tuple[int, dict]]  # (1-based source row number, {header: value})
    sheet_name: str | None = None
    header_row: int = 1
    sheet_names: list[str] = field(default_factory=list)


# extension -> reader(content, sheet, header_row, sheet_hint) -> TableData
# Word/PDF extractors plug in here later: they must return the same TableData (canonical-header rows) so that
# everything downstream (mapping, validation, loaders, exceptions, KPIs) is shared with Excel/CSV.
READERS: dict[str, Callable[..., "TableData"]] = {}


def register_reader(extensions: tuple[str, ...], fn: Callable[..., "TableData"]) -> None:
    for ext in extensions:
        READERS[ext] = fn


def read_table(filename: str, content: bytes, sheet: str | None = None, header_row: int | None = None,
               sheet_hint: str | None = None) -> TableData:
    ext = PurePath(filename).suffix.lower()
    reader = READERS.get(ext)
    if reader is None:
        raise UnsupportedFileError(
            f"No reader configured for '{ext}'. Supported: {', '.join(sorted(READERS))}. "
            "Word/PDF documents need an extractor (not configured yet).")
    return reader(content, sheet, header_row, sheet_hint)


def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "cp1256", "latin-1"):  # cp1256 = Windows Arabic
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    raise UnsupportedFileError("Could not decode file text")


def _read_csv(content: bytes, sheet: str | None, header_row: int | None, sheet_hint: str | None = None) -> TableData:
    text = _decode(content)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    raw = [tuple(r) for r in csv.reader(io.StringIO(text), dialect)]
    return _build_table(raw, header_row, sheet_name=None, sheet_names=[])


def _read_xlsx(content: bytes, sheet: str | None, header_row: int | None, sheet_hint: str | None = None) -> TableData:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        names = wb.sheetnames
        if sheet is not None and sheet not in names:
            raise UnsupportedFileError(f"Sheet '{sheet}' not found. Available: {names}")
        if sheet is None and sheet_hint is not None:
            if sheet_hint in names:
                sheet = sheet_hint
            elif len(names) > 1:  # never silently read a different sheet of a multi-sheet workbook
                raise UnsupportedFileError(f"Expected sheet '{sheet_hint}'. Available: {names}. Pass sheet=...")
        ws = wb[sheet] if sheet else wb[names[0]]
        raw = [tuple(r) for r in ws.iter_rows(values_only=True)]
        return _build_table(raw, header_row, sheet_name=ws.title, sheet_names=names)
    finally:
        wb.close()


def _detect_header_row(raw: list[tuple]) -> int:
    """First of the top rows that is mostly text cells and about as wide as the widest row."""
    scan = raw[:HEADER_SCAN_ROWS]
    widest = max((sum(c is not None and str(c).strip() != "" for c in r) for r in scan), default=0)
    for idx, row in enumerate(scan):
        filled = [c for c in row if c is not None and str(c).strip() != ""]
        texty = [c for c in filled if isinstance(c, str)]
        if filled and len(filled) >= max(2, 0.6 * widest) and len(texty) >= 0.8 * len(filled):
            return idx + 1
    return 1


def _build_table(raw: list[tuple], header_row: int | None, sheet_name: str | None, sheet_names: list[str]) -> TableData:
    if not raw:
        return TableData([], [], sheet_name, 1, sheet_names)
    hdr = header_row or _detect_header_row(raw)
    header_cells = raw[hdr - 1]
    headers: list[str] = []
    for i, cell in enumerate(header_cells):
        name = clean_text(cell) or f"column_{i + 1}"
        while name in headers:  # keep duplicate headers addressable
            name += "_2"
        headers.append(name)
    rows = []
    for offset, row in enumerate(raw[hdr:], start=hdr + 1):
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        padded = list(row) + [None] * (len(headers) - len(row))
        rows.append((offset, dict(zip(headers, padded[: len(headers)]))))
    return TableData(headers, rows, sheet_name, hdr, sheet_names)


register_reader((".xlsx", ".xlsm"), _read_xlsx)
register_reader((".csv", ".txt"), _read_csv)
