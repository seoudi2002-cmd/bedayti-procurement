"""Excel exporter (openpyxl): summary sheet, one sheet per table (real numbers with number formats, filters, frozen
header) and a charts sheet with the same images as the PDF. Sheets are right-to-left for Arabic reports."""
import io
import re
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.reporting.base import ReportModel
from app.core.reporting.charts import render_chart

FORMATS = {"money": "#,##0.00", "smoney": "+#,##0.00;-#,##0.00;0.00", "pct": '0.0"%"', "spct": '+0.0"%";-0.0"%"', "int": "#,##0"}
HEAD = PatternFill("solid", fgColor="2A78D6")
CARD = PatternFill("solid", fgColor="F4F6FA")


def _name(title: str, used: set) -> str:
    base = re.sub(r"[\[\]\*\?/\\:]", " ", title).strip()[:28] or "Sheet"
    n, i = base, 2
    while n.lower() in used:
        n = f"{base[:25]} {i}"
        i += 1
    used.add(n.lower())
    return n


def _val(v):
    if isinstance(v, Decimal):
        return float(v)
    return v


class ExcelExporter:
    extension = "xlsx"

    def render(self, report: ReportModel) -> bytes:
        rtl = report.lang == "ar"
        wb = Workbook()
        ws = wb.active
        used: set = set()
        ws.title = _name(report.sections[0].title if report.sections else "Summary", used)
        ws.sheet_view.rightToLeft = rtl
        ws["A1"] = report.title
        ws["A1"].font = Font(bold=True, size=15)
        ws["A2"] = report.subtitle or ""
        ws["A3"] = report.period_label
        r = 5
        for sec in report.sections[:1]:
            for k in sec.kpis:
                ws.cell(r, 1, k["label"]).font = Font(color="52514E", size=9)
                c = ws.cell(r, 2, k["value"])
                c.font = Font(bold=True, size=12)
                ws.cell(r, 3, k.get("sub", ""))
                for col in (1, 2, 3):
                    ws.cell(r, col).fill = CARD
                r += 1
            r += 1
            for i in sec.insights:
                ws.cell(r, 1, "• " + i["text"]).alignment = Alignment(wrap_text=True, vertical="top")
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
                ws.row_dimensions[r].height = 15 * (1 + len(i["text"]) // 110)
                r += 1
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 26
        ws.column_dimensions["C"].width = 24
        for sec in report.sections:
            for tb in sec.tables:
                if not tb["rows"]:
                    continue
                sh = wb.create_sheet(_name(tb["title"], used))
                sh.sheet_view.rightToLeft = rtl
                sh["A1"] = tb["title"]
                sh["A1"].font = Font(bold=True, size=12)
                for j, c in enumerate(tb["columns"], 1):
                    cell = sh.cell(3, j, c["label"])
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = HEAD
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                for i, row in enumerate(tb["rows"], 4):
                    for j, c in enumerate(tb["columns"], 1):
                        v = _val(row.get(c["key"]))
                        cell = sh.cell(i, j, v)
                        if c["fmt"] in FORMATS and isinstance(v, (int, float)):
                            cell.number_format = FORMATS[c["fmt"]]
                        elif isinstance(v, str):
                            cell.alignment = Alignment(wrap_text=True, vertical="top")
                for j, c in enumerate(tb["columns"], 1):
                    width = max([len(str(c["label"]))] + [len(str(row.get(c["key"]) or "")) for row in tb["rows"][:200]])
                    sh.column_dimensions[get_column_letter(j)].width = min(max(width + 2, 10), 60 if c["fmt"] == "text" else 18)
                sh.freeze_panes = "A4"
                sh.auto_filter.ref = f"A3:{get_column_letter(len(tb['columns']))}{3 + len(tb['rows'])}"
        charts = [ch for sec in report.sections for ch in sec.charts]
        if charts:
            cs = wb.create_sheet(_name("Charts" if not rtl else "الرسوم", used))
            cs.sheet_view.rightToLeft = rtl
            row = 1
            for ch in charts:
                try:
                    png = render_chart(ch, report.lang)
                except Exception:
                    continue
                img = XlImage(io.BytesIO(png))
                scale = 620 / img.width
                img.width, img.height = 620, int(img.height * scale)
                cs.add_image(img, f"A{row}")
                row += int(img.height / 20) + 3
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
