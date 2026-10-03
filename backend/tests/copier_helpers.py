"""Synthetic copier files that mimic the structure of the real ones (no real data)."""
import functools
import io

from openpyxl import Workbook

H1 = ["م", None, "الفرع", "القراءة\nالسابقة", "القراءة\nالحالية", "الاستهلاك", "الاضافى"]
H2 = ["م", None, "الفرع", "الماكينة", "القراءة السابقة", "القراءة الحالية", "الاستهلاك", "الاضافى", "قيمه الايجاريه", "ملاحظات"]


def _put(ws, r, values, start=2):
    for i, v in enumerate(values):
        if v is not None:
            ws.cell(r, start + i, v)


@functools.lru_cache(maxsize=None)  # saved workbooks/PDFs embed a timestamp: identical bytes are needed for duplicate-upload tests
def statement_workbook(month: int = 7, shift: int = 0, drop_part2_row: bool = False, break_total: bool = False) -> bytes:
    """Two classes (copiers 5000: 3 machines, printers 3000: 2 machines) + Part 2 by governorate. `shift` moves all readings
    forward (to build a second month whose previous readings equal the first month's current ones)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["D1"] = f"زيارات شركه بدايتى {month} - 2026   ماكينات - 5000 نسخه"
    _put(ws, 3, H1)
    m5 = [("فرع ألف B", 1000, 4000), ("فرع باء B", 2000, 8000), ("فرع جيم", 3000, 3500)]   # (branch, prev, consumption)
    r = 4
    for i, (b, prev, cons) in enumerate(m5, 1):
        prev, cur = prev + shift, prev + shift + cons
        _put(ws, r, [i, None, b, prev, cur, cur - prev, max(0, cons - 5000)])
        r += 1
    ws.cell(r, 2, "الاجمالى")
    ws.cell(r, 8, 3000 if not break_total else 9999)
    ws["D12"] = f"زيارات شركه بدايتى {month}- 2026   طابعات - 3000 نسخه"
    _put(ws, 14, H1)
    p3 = [("فرع ألف", 500, 200), ("فرع باء", 900, 3600)]
    r = 15
    for i, (b, prev, cons) in enumerate(p3, 1):
        prev, cur = prev + shift, prev + shift + cons
        _put(ws, r, [i, None, b, prev, cur, cur - prev, max(0, cons - 3000)])
        r += 1
    ws.cell(r, 2, "الاجمالى")
    ws.cell(r, 8, 600)
    # ---- Part 2: by governorate (supporting)
    ws["D20"] = "الشرقيه"
    _put(ws, 21, H2)
    rows = [("فرع ألف", 5000, 1000 + shift, 5000 + shift, 4000, 0, 1045), ("فرع ألف", 3000, 500 + shift, 700 + shift, 200, 0, 660),
            ("فرع باء", 5000, 2000 + shift, 10000 + shift, 8000, 3000, 1045), ("فرع باء", 3000, 900 + shift, 4500 + shift, 3600, 600, 660)]
    r = 22
    for br, mach, prev, cur, cons, exc, rent in rows:
        if drop_part2_row and br == "فرع ألف" and mach == 3000:
            continue
        _put(ws, r, [None, None, br, mach, prev, cur, cons, exc, rent])
        r += 1
    ws.cell(r + 1, 4, "قنا")
    _put(ws, r + 2, H2)
    _put(ws, r + 3, [None, None, "فرع جيم", 5000, 3000 + shift, 6500 + shift, 3500, 0, 1045])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@functools.lru_cache(maxsize=None)
def invoice_pdf(qty_5000: int = 3, qty_3000: int = 2, exc_5000: int = 3000, exc_3000: int = 600, sales_wrong: bool = False) -> bytes:
    """A PDF with the layout features that matter to the reader: visual-order Arabic (presentation forms), a 12-column
    header table, item 'EGS' codes, numbers split across lines, a tax sub-table and a totals block."""
    import arabic_reshaper
    from bidi.algorithm import get_display
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    import matplotlib as mpl
    fonts = mpl.get_data_path() + "/fonts/ttf/DejaVuSans.ttf"
    pdfmetrics.registerFont(TTFont("DJ", fonts))

    def vis(t):  # visual order, like the e-invoice PDF
        return get_display(arabic_reshaper.reshape(t))
    buf = io.BytesIO()
    W, H = 760, A4[1]
    c = canvas.Canvas(buf, pagesize=(W, H))
    c.setFont("DJ", 8)
    c.drawString(40, H - 30, "eInvoicing")
    lines = [("الرقم الداخلى:", None), ("1976", None), ("تاريخ الإصدار:", None), ("RAW:م 6:34:51 2026/ 8/ 9", None), ("حالة الفاتورة:", None), ("صحيح", None),
             ("رقم التسجيل:", None), ("631284974", None), ("رقم التسجيل للبائع:", None), ("489283268", None)]
    y = H - 60
    c.drawString(300, y, "2BYHN2XAHNQ994SQTHQKJKZK10")
    y -= 14
    for t, _ in lines:
        c.drawRightString(W - 40, y, t[4:] if t.startswith("RAW:") else (vis(t) if not t[0].isdigit() else t))
        y -= 14
    c.drawRightString(W - 40, y - 6, vis("الإسم"))
    c.drawRightString(W - 330, y - 6, vis("الإسم"))
    c.drawRightString(W - 40, y - 24, vis("شركة المورد التجريبية"))
    c.drawRightString(W - 330, y - 24, vis("بدايتي لتمويل المشروعات"))
    # header table: 12 columns, right to left roles
    x0 = 20
    cols = [(0, "المجموع / خصم الأصناف"), (1, "إجمالي الرسوم الخاضعة للضريبة"), (2, "صافي الإجمالي"), (3, "قيمة فرق لأغراض الضريبة"), (4, "قيمة الضريبة"),
            (5, "قيمة الخصم / نسبة الخصم"), (6, "المبلغ الإجمالي للمبيعات"), (7, "سعر الوحدة"), (8, "نوع / الكمية الوحدة"), (9, "الكود الداخلي / كود الصنف"),
            (10, "مخطط نظام التكويد"), (11, "اسم الكود/الوصف")]
    widths = [44, 44, 40, 40, 36, 40, 40, 36, 40, 40, 36, 260]
    xs = [x0]
    for w in widths:
        xs.append(xs[-1] + w)
    top = 520
    c.rect(xs[0], top, xs[-1] - xs[0], 36)
    for i, (_, label) in enumerate(cols):
        c.line(xs[i], top, xs[i], top + 36)
        c.setFont("DJ", 4.5)
        words, cur, wrapped = label.split(), "", []
        for w_ in words:  # wrap by word so every label keeps its keywords
            if cur and len(cur) + len(w_) > 11:
                wrapped.append(cur)
                cur = w_
            else:
                cur = f"{cur} {w_}".strip()
        wrapped.append(cur)
        for k, line in enumerate(wrapped[:4]):
            c.drawCentredString((xs[i] + xs[i + 1]) / 2, top + 28 - 6 * k, vis(line))
    c.line(xs[-1], top, xs[-1], top + 36)

    def item(ytop, desc_lines, qty, unit, sales):
        vat, wht = sales * 0.14, sales * 0.03
        total = sales + vat - wht
        c.setFont("DJ", 6)
        for k, d in enumerate(desc_lines):
            c.drawRightString(xs[-1] - 3, ytop - 8 * k, vis(d) if not d[0].isdigit() else d)
        c.drawString(xs[10] + 6, ytop - 10, "EGS")
        c.drawString(xs[8] + 4, ytop - 10, f"{qty:,.4f}")
        c.drawString(xs[8] + 4, ytop - 18, "JOB")
        c.drawString(xs[7] + 4, ytop - 10, f"{unit:,.3f}")
        c.drawString(xs[7] + 4, ytop - 18, "00")
        c.drawString(xs[6] + 3, ytop - 10, f"{sales:,.2f}")
        c.drawString(xs[6] + 3, ytop - 18, "000")
        c.drawString(xs[2] + 3, ytop - 10, f"{sales:,.2f}")
        c.drawString(xs[2] + 3, ytop - 18, "000")
        c.drawString(xs[0] + 3, ytop - 10, f"{total:,.2f}")
        c.drawString(xs[0] + 3, ytop - 18, "000")
        # tax sub-table: values in the tax column, rates in the sales column
        c.drawString(xs[4] + 2, ytop - 44, f"{vat:,.5f}")
        c.drawString(xs[6] + 3, ytop - 44, "14.00%")
        c.drawString(xs[4] + 2, ytop - 54, f"{wht:,.5f}")
        c.drawString(xs[6] + 3, ytop - 54, "3.00%")
        return sales
    s1 = item(top - 20, ["تأجير ماكينات تصوير مستندات/ ماكينات تصوير مستندات", "5000 نسخه في الفترة من 2026-7-1 الى 2026-7-31"], qty_5000, 1045, qty_5000 * 1045)
    s2 = item(top - 110, ["تأجير ماكينات تصوير مستندات/ تأجير طابعات مستندات", "3000 نسخه في الفترة من 2026-7-1 الى 2026-7-31"], qty_3000, 660, qty_3000 * 660)
    s3 = item(top - 200, ["تأجير ماكينات تصوير مستندات/ اضافى النسخ لماكينات", "5000 نسخه في الفترة من 2026-7-1 الى 2026-7-31"], exc_5000, 0.319, exc_5000 * 0.319)
    s4 = item(top - 290, ["تأجير ماكينات تصوير مستندات/ اضافي النسخ طابعات", "في الفترة من 2026-7-1 الى 2026-7-31"], exc_3000, 0.319, exc_3000 * 0.319)
    sales = s1 + s2 + s3 + s4
    if sales_wrong:
        sales += 10
    vat, wht = sales * 0.14, sales * 0.03
    y = 120
    for val, label in [(sales, "(م.ج) إجمالي المبيعات"), (0, "(م.ج) إجمالي الخصم"), (0, "(م.ج) إجمالي خصم الأصناف"), (vat, "(م.ج) ضريبة القيمة المضافة"),
                       (wht, "(م.ج) الخصم تحت حساب الضريبة"), (0, "(م.ج) خصم إضافي على إجمالي الفاتورة"), (sales + vat - wht, "(م.ج) المبلغ الإجمالي")]:
        c.rect(40, y, 90, 12)
        c.rect(130, y, 160, 12)
        c.setFont("DJ", 6)
        c.drawString(44, y + 3, f"{val:,.5f}")
        c.drawRightString(286, y + 3, vis(label))
        y -= 12
    c.save()
    return buf.getvalue()


@functools.lru_cache(maxsize=None)
def paper_workbook(po: str = "35", received: int | None = 20, break_month: bool = False, with_dup: bool = True) -> bytes:
    """Paper distribution statement (synthetic): title with PO/receipt, rows (م, عدد, فرع, تاريخ), monthly subtotals in E/F and a total."""
    from datetime import datetime
    wb = Workbook()
    ws = wb.active
    ws.title = "الاستهلاك"
    ws["A1"] = f"بيان توزيع ورق التصوير والطباعة أمر شراء رقم ({po}) من تاريخ 20 مايو 2026 (استلام {received} كرتونة)" if received is not None \
        else f"بيان توزيع ورق التصوير والطباعة أمر شراء رقم ({po}) من تاريخ 20 مايو 2026"
    for c, v in enumerate(["م", "العدد", "الفرع", "التاريخ", "العدد", "البيان - استلام / متبقي"], 1):
        ws.cell(2, c, v)
    rows = [(1, 4, "فرع ألف", datetime(2026, 7, 2)), (2, 4, "فرع باء", datetime(2026, 7, 2)), (3, 2, "المركز الرئيسي - الحفظ", datetime(2026, 7, 5)),
            (4, 3, "المركز الرئيسي الحفظ", datetime(2026, 7, 9)), (5, 1, "المركز الرئيسي", datetime(2026, 7, 9)),
            (7, 4, "فرع ألف", datetime(2026, 8, 3)), (8, 2, "المركز الرئيسي - HR", datetime(2026, 8, 3)), (9, 0.5, "فرع جيم", datetime(2026, 8, 10))]
    if with_dup:
        rows.insert(2, (10, 4, "فرع باء", datetime(2026, 7, 2)))      # same branch, same day, a second line
    r = 3
    for seq, cartons, br, d in rows:
        ws.cell(r, 1, seq)
        ws.cell(r, 2, cartons)
        ws.cell(r, 3, br)
        ws.cell(r, 4, d)
        r += 1
    jul = sum(x[1] for x in rows if x[3].month == 7)
    aug = sum(x[1] for x in rows if x[3].month == 8)
    ws.cell(5, 5, jul + (1 if break_month else 0))
    ws.cell(5, 6, "استهلاك شهر يوليو")
    ws.cell(r - 1, 5, aug)
    ws.cell(r - 1, 6, "استهلاك شهر اغسطس")
    ws.cell(r + 1, 1, "العدد الاجمالى")
    ws.cell(r + 1, 2, jul + aug)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
