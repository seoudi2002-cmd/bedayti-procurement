"""Rent, vehicles and overtime: time-varying operating data kept as versions (nothing replaced, differences shown)."""
import pytest
from sqlalchemy import func, select

from app.core.analysis import AnalysisError
from app.core.reporting.excel import ExcelExporter
from app.core.reporting.pdf import PdfExporter
from app.core.settings_store import effective_thresholds
from app.models import OpRecord
from app.modules.overtime_analysis import engine as ot_engine
from app.modules.overtime_analysis import service as ot_service
from app.modules.overtime_analysis import workbook as ot_wb
from app.modules.overtime_analysis.adapter import OvertimeAdapter
from app.modules.rent_analysis import engine as rent_engine
from app.modules.rent_analysis import service as rent_service
from app.modules.rent_analysis import workbook as rent_wb
from app.modules.rent_analysis.adapter import RentAdapter
from app.modules.vehicle_analysis import cards, grid, plates, repairs
from app.modules.vehicle_analysis import service as veh_service
from app.modules.vehicle_analysis.adapter import VehicleAdapter
from tests.ops_helpers import card_book, overtime_workbook, rent_workbook, repairs_workbook, usage_book
from tests.test_copier_analysis import api, auth  # noqa: F401


def issue(parsed, code):
    return next((i for i in parsed.issues.items if i.code == code), None)


# ================================================================================================ rent
def test_rent_reader_takes_the_governorate_from_explicit_evidence_and_ignores_stale_copies():
    p = rent_wb.parse(rent_workbook(1))
    a = rent_wb.assemble(p)
    by = {c["copy"].name: c for c in a["contracts"]}
    assert len(by) == 6
    assert (by["الزقاريق ( الشرقية )"]["governorate"], by["الزقاريق ( الشرقية )"]["evidence"]) == ("الشرقية", "sheet_formula")
    assert by["بنها ( المنوفية )"]["evidence"] == "sheet_formula" and by["بنها ( المنوفية )"]["copy"].footer        # the sheet's own contract below its master list
    assert by["العجوزة د3"]["governorate"] == "المركز الرئيسي"
    assert by["فرع بلا محافظة"]["governorate"] is None and issue(p, "governorate_unresolved")                       # no evidence -> not allocated
    # the Menoufia sheet holds a stale copy of a Sharqia contract (999): the owning sheet's value is the one used
    assert by["الزقاريق ( الشرقية )"]["copy"].months["2025-03"] == 1100 and issue(p, "stale_copies")
    assert by["أبو حماد ( الشرقية )"]["copy"].nopay == {"2025-02"} and "2025-02" not in by["أبو حماد ( الشرقية )"]["copy"].months   # «_____» is not zero
    assert issue(p, "end_unreadable").examples == ["فرع بلا محافظة: 31/9/2030"]                                    # an impossible date is reported, not repaired
    assert issue(p, "month_header_year_inconsistent") and by["العجوزة د3"]["copy"].months["2025-03"] == 770     # a header typo is placed by sequence and reported


def test_rent_versions_keep_the_history_and_show_what_changed(session):
    ad = RentAdapter()
    ad.ingest(session, rent_workbook(1), "rent_a.xlsx", "bob", {})
    n1 = session.scalar(select(func.count()).select_from(OpRecord))
    with pytest.raises(AnalysisError) as e:
        ad.ingest(session, rent_workbook(1), "again.xlsx", "bob", {})
    assert e.value.status == 409
    ad.ingest(session, rent_workbook(2), "rent_b.xlsx", "bob", {})
    assert session.scalar(select(func.count()).select_from(OpRecord)) > n1                      # the second file adds records; the first file's are still there
    d = rent_service.load(session)
    ch = [c for c in d["changes"] if c["kind"] == "rent_value"]
    assert len(ch) == 1 and (ch[0]["period"], ch[0]["field"], ch[0]["old"], ch[0]["new"]) == ("2025-03", "rent", 1100, 1210)
    a = next(c for c in d["contracts"] if c["name"].startswith("الزقاريق"))
    assert a["months"]["2025-03"] == 1210 and a["versions"] == 2                                  # effective value = newest file; the old one is kept
    old = session.scalars(select(OpRecord).where(OpRecord.kind == "rent_value", OpRecord.period == "2025-03", OpRecord.dataset_id == 1)).all()
    assert 1100 in [r.values["rent"] for r in old]
    assert [v["changed"] for v in d["versions"]] == [0, 1]
    with pytest.raises(AnalysisError) as e:
        ad.delete_item(session, "all")
    assert e.value.status == 409


def test_rent_analysis_totals_by_month_and_governorate_and_the_files_own_subtotals(session):
    ad = RentAdapter()
    ad.ingest(session, rent_workbook(1), "rent_a.xlsx", "bob", {})
    rm, a = ad.build(session, "all", "en", True, {})
    assert [round(m["total"]) for m in a["months"]] == [4880, 2680, 5390]                            # Feb: contract B has no payment recorded (not 0, not counted)
    assert a["months"][1]["nopay"] == 1
    g = {x["name"]: x for x in a["governorates"]}
    assert g["الشرقية"]["last"] == 3300 and g["المنوفية"]["last"] == 990 and g[None]["last"] == 330 and g["المركز الرئيسي"]["last"] == 770
    assert a["totals"]["unresolved"] == 1 and a["totals"]["expiring"] == 1                         # B ends 2025-03-31 — inside the window? it ends in the last month
    ctl = {c["sheet"]: c for c in rent_service.load(session)["summaries"][0]["controls"]}
    assert (ctl["الشرقية"]["matched"], ctl["الشرقية"]["mismatched"]) == (3, 0)
    assert (ctl["المنوفية"]["matched"], ctl["المنوفية"]["mismatched"]) == (2, 1) and ctl["المنوفية"]["examples"][0]["stated"] == 991
    assert ctl["المركز الرئيسي"]["matched"] == 3
    keys = [s.key for s in rm.sections]
    assert keys == ["summary", "monthly", "governorates", "contracts", "steps", "expiry", "versions", "quality"]
    # landlord names (personal data) only for admins
    cols_admin = [c["key"] for t in rm.sections[3].tables for c in t["columns"]]
    cols_viewer = [c["key"] for t in ad.build(session, "all", "en", False, {})[0].sections[3].tables for c in t["columns"]]
    assert "ll" in cols_admin and "ll" not in cols_viewer
    snap, sa = ad.build(session, "m:2025-02", "ar", False, {})
    assert sa["totals"]["last"] == "2025-02" and round(sa["totals"]["total_last"]) == 2680
    assert PdfExporter().render(rm)[:4] == b"%PDF" and ExcelExporter().render(rm)[:2] == b"PK"
    with pytest.raises(AnalysisError):
        ad.build(session, "all", "en", True, {"governorates": ["لا توجد"]})


def test_rent_step_detection_uses_recorded_values_only():
    contracts = [{"key": "k", "name": "x", "governorate": "g", "scope": "branch", "start": None, "end": None, "start_raw": None, "end_raw": None, "advance": None, "deposit": None, "contract_rent": None,
                  "current_rent_stated": None, "landlord": None, "flags": [], "months": {"2025-01": 100.0, "2025-02": 110.0, "2025-04": 121.0, "2025-05": 300.0}, "nopay": {"2025-03"}, "versions": 1}]
    th = effective_thresholds.__globals__["default_thresholds"]("rent")
    a = rent_engine.analyze(contracts, th)
    assert [(s["period"], round(s["pct"])) for s in reversed(a["steps"])] == [("2025-02", 10), ("2025-04", 10), ("2025-05", 148)]
    assert [s["large"] for s in reversed(a["steps"])] == [False, False, True]                         # a jump above the threshold is listed as large, not as a regular increase
    assert a["totals"]["periodic"] == 0 and a["rows"][0]["blank_months"] == 1


# ================================================================================================ overtime
def test_overtime_reads_the_period_from_the_title_and_excludes_a_repeated_sheet():
    p = ot_wb.parse(overtime_workbook())
    sheets = ot_wb.select(p)
    ot_wb.check(p, sheets)
    assert [(s.sheet, s.period) for s in sheets] == [("يناير", "2025-01"), ("فبراير", "2025-02")]
    assert issue(p, "duplicate_period_sheet").examples == ["مارس = فبراير (2025-02)"] and issue(p, "sheet_name_title_mismatch")
    assert p.factors == (1.35, 1.7) and not issue(p, "weighted_differs_from_multiplier")             # the multipliers are the file's own
    assert [r.code for r in sheets[0].rows] == ["10", "11", "12"] and sheets[0].rows[2].values["day_hours"] is None  # listed, no entry


def test_overtime_ingest_reconciles_the_annual_sheet_and_keeps_versions(session):
    ad = OvertimeAdapter()
    r = ad.ingest(session, overtime_workbook(), "ot.xlsx", "bob", {})
    assert r.meta["months"] == ["2025-01", "2025-02"]
    ctl = {c["field"]: c for c in ot_service.load(session)["summaries"][0]["controls"]}
    assert ctl["day_hours"]["stated"] == 130 and ctl["day_hours"]["computed"] == 80 and ctl["day_hours"]["diff"] == 50    # the annual sheet counts the repeated month
    ad.ingest(session, overtime_workbook(True), "ot_corrected.xlsx", "bob", {})
    d = ot_service.load(session)
    ch = {c["field"]: (c["old"], c["new"]) for c in d["changes"]}
    assert ch["day_hours"] == (40, 44) and ch["day_weighted"] == (54, 59.4) and {c["period"] for c in d["changes"]} == {"2025-02"}
    feb10 = next(x for x in d["rows"] if x["code"] == "10" and x["period"] == "2025-02")
    assert feb10["day_hours"] == 44 and feb10["versions"] == 2
    assert session.scalar(select(func.count()).select_from(OpRecord).where(OpRecord.kind == "overtime", OpRecord.entity_key == "10", OpRecord.period == "2025-02")) == 2   # both versions stored


def test_overtime_analysis_hours_only_no_money_and_names_for_admins_only(session):
    ad = OvertimeAdapter()
    ad.ingest(session, overtime_workbook(), "ot.xlsx", "bob", {})
    rm, a = ad.build(session, "all", "en", True, {})
    m = a["avail"]
    assert [(x["hours"], round(x["weighted"], 1), x["listed"], x["entries"], x["active"]) for x in m] == [(45.0, 66.0, 3, 2, 2), (62.0, 87.9, 3, 2, 2)]
    assert a["totals"]["empty_rows"] == 2                                                              # code 12: listed, no entry (not zero)
    e = {x["code"]: x for x in a["employees"]}
    assert e["10"]["hours"] == 82 and e["12"]["n"] == 0 and e["12"]["avg_month"] is None
    text = " ".join(x["text"] for sec in rm.sections for x in sec.insights)
    assert "ج.م" not in text and "EGP" not in text                                                    # nothing is converted to money
    tbl = lambda r: next(t for sec in r.sections if sec.key == "employees" for t in sec.tables if t["key"] == "employees")
    assert "name" in [c["key"] for c in tbl(rm)["columns"]]
    assert "name" not in [c["key"] for c in tbl(ad.build(session, "all", "en", False, {})[0])["columns"]]
    assert "موظف أول" not in repr(ad.build(session, "all", "en", False, {})[0].sections[2])
    snap, sa = ad.build(session, "m:2025-01", "ar", False, {})
    assert sa["totals"]["ytd_months"] == 1 and sa["totals"]["last"] == "2025-01"
    assert PdfExporter().render(rm)[:4] == b"%PDF" and ExcelExporter().render(rm)[:2] == b"PK"


def test_overtime_gaps_are_not_available_not_zero():
    rows = [{"code": "1", "name": "a", "period": p, "day_hours": h, "night_hours": 0.0, "day_weighted": None, "night_weighted": None, "weighted_total": None, "mission_day": None, "mission_night": None,
             "meals": None, "equal_pay_days": None, "equal_pay_x2": None, "raw_total": None, "versions": 1} for p, h in (("2025-01", 10.0), ("2025-03", 30.0))]
    roster = [{"code": "1", "name": "a", "period": r["period"], "versions": 1} for r in rows]
    a = ot_engine.analyze(rows, roster, effective_thresholds.__globals__["default_thresholds"]("overtime"))
    assert a["totals"]["missing"] == ["2025-02"] and [m["available"] for m in a["months"]] == [True, False, True]
    assert a["months"][2].get("gap_before") is True and a["months"][2]["large"] is False             # no month-on-month comparison across a missing month


# ================================================================================================ vehicles
def test_repairs_reader_keeps_stated_totals_and_compares_them_with_their_parts():
    gs = grid.grids(repairs_workbook(), "repairs.xlsx")
    assert repairs.is_repairs(gs)
    p = repairs.parse(gs, None, "repairs.xlsx")
    assert (p.year, p.year_source, p.periods) == (2026, "other_sheets", ["2026-01", "2026-02"])
    r = next(x for x in p.rows if x.plate == "س ص 1111" and x.period == "2026-01")
    assert r.values["grand_total"] == 1120 and r.values["cat:بوابات رسوم"] == 20 and r.values["fuel_price_in_formula"] == 20.0 and "fuel_qty" not in r.values   # price read from the formula; no value is invented
    assert {c["period"]: c["diff"] for c in p.controls} == {"2026-01": 1.0, "2026-02": None}                       # a stated total row one off; a month with no total row
    assert issue(p, "total_row_differs") and issue(p, "no_total_row")
    half = {h["plate"]: h["diff"] for h in p.halfyear}
    assert half["س ص 1111"] == 7069 and half["س ص 2222"] == 0 and half[None] == 19499 - 12430 and issue(p, "halfyear_total_differs")
    assert [(c["desc"], c["period"], c["values"]["company_share"]) for c in p.claims] == [("زجاج", "2026-01", 300)]


def test_usage_report_and_card_readers_keep_stated_figures_and_skip_unfilled_months():
    u = cards.parse_usage(grid.grids(usage_book(), "u.xlsx"), "u.xlsx")
    assert [(x.period, x.plate, x.values["km_total"], x.values["days_used"], x.values["odometer_last"]) for x in u.rows] == [("2026-01", "ع ن 3333", 150, 2, 1150), ("2026-02", "ع ن 3333", 150, 2, 1300)]
    assert issue(u, "sheet_name_date_mismatch")                                                       # the date cell moves by a day, the sheet name carries the month
    c = cards.parse_card(grid.grids(card_book(), "c.xlsx"), "c.xlsx")
    assert [(x.period, x.plate, x.values["odometer"]) for x in c.rows] == [("2026-01", "3333", 29500), ("2026-02", "3333", 29900)]
    assert c.skipped == ["2026-03"] and issue(c, "template_months_not_recorded")                       # odometer 0 = unfilled template, not a reading


def test_plates_link_only_by_exact_digits_only_unique_or_approved_alias():
    keys = [plates.key("ع م 3333"), plates.key("ع ن 3333"), plates.key(3333.0), plates.key("س ص 1111")]
    canon, cand = plates.link(keys)
    assert canon[plates.key("ع م 3333")] != canon[plates.key("ع ن 3333")]                          # a different letter is not guessed away
    assert canon[plates.key(3333.0)] == plates.key(3333.0) and cand[0][0] == "3333"                  # two lettered plates carry the number: the digits-only source stays apart
    canon, _c = plates.link(keys, {"ع ن 3333": "ع م 3333"})                                         # an approved alias links them, and then the number is unique
    assert canon[plates.key("ع ن 3333")] == canon[plates.key(3333.0)] == plates.key("ع م 3333")
    canon, cand = plates.link([plates.key("ع ن 3333"), plates.key(3333.0)])
    assert canon[plates.key(3333.0)] == plates.key("ع ن 3333") and not cand


def test_vehicle_versions_analysis_and_linking(session):
    ad = VehicleAdapter()
    ad.ingest(session, repairs_workbook(), "repairs.xlsx", "bob", {})
    ad.ingest(session, usage_book(), "usage.xlsx", "bob", {})
    ad.ingest(session, card_book(), "card.xlsx", "bob", {})
    d = veh_service.load(session)
    assert sorted(v["plate"] for v in d["vehicles"]) == sorted(["س ص 1111", "ع ن 3333", "س ص 2222"])          # the card's digits-only «3333» joined the one lettered plate with that number
    v3333 = next(v for v in d["vehicles"] if v["plate"] == "ع ن 3333")
    assert sorted(v3333["usage"]) == ["2026-01", "2026-02"] and sorted(v3333["service"]) == ["2026-01", "2026-02"] and not v3333["cost"]
    rm, a = ad.build(session, "all", "en", True, {})
    assert [round(m["total"]) for m in a["months"]] == [8120, 4310] and round(a["totals"]["maint"]) == 6230 and round(a["totals"]["fuel"]) == 6200
    veh = {v["plate"]: v for v in a["vehicles"]}
    assert veh["س ص 2222"]["total"] == 9500 and round(veh["س ص 2222"]["cost_per_km"], 3) == round(9500 / 3100, 3)
    assert round(a["totals"]["cost_per_km"], 3) == round(12430 / 4800, 3)                              # only vehicle-months that have a distance
    assert veh["ع ن 3333"]["total"] == 0 and veh["ع ن 3333"]["months"] == 0                       # no cost figure for it: not a zero-cost car
    assert [(x["plate"], x["state"], x["remaining"]) for x in a["due"]] == [("ع ن 3333", "soon", 100)]
    assert a["totals"]["claims"] == 1 and a["totals"]["claims_company"] == 300
    text = " ".join(x["text"] for sec in rm.sections for x in sec.insights)
    assert "derived" in text                                                                          # the fuel quantity is labelled as derived, not measured
    # a corrected repairs file is a new version: the first is kept and the difference is listed
    ad.ingest(session, repairs_workbook(True), "repairs_corrected.xlsx", "bob", {})
    d2 = veh_service.load(session)
    ch = {c["field"]: (c["old"], c["new"]) for c in d2["changes"] if c["period"] == "2026-02"}
    assert ch == {"cat:صيانة": (700, 900), "maint_total": (710, 910), "grand_total": (1810, 2010)}
    assert [v["changed"] for v in d2["versions"] if v["layout"] == "repairs_statement"] == [0, 3]
    rm2, a2 = ad.build(session, "m:2026-01", "ar", False, {"vehicles": [v["key"] for v in d2["vehicles"] if v["plate"] == "س ص 1111"]})
    assert a2["totals"]["total"] == 1120 and a2["filtered"]
    assert PdfExporter().render(rm).startswith(b"%PDF") and ExcelExporter().render(rm)[:2] == b"PK"


# ================================================================================================ API
def test_generic_api_for_the_operating_modules(api):
    for key, files, item in (("rent", [("rent.xlsx", rent_workbook(1)), ("rent2.xlsx", rent_workbook(2))], "m:2025-03"),
                             ("overtime", [("ot.xlsx", overtime_workbook())], "all"),
                             ("vehicles", [("repairs.xlsx", repairs_workbook()), ("usage.xlsx", usage_book())], "all")):
        base = f"/api/analysis/{key}/datasets"
        assert api.post(base, files={"file": (files[0][0], files[0][1])}, headers=auth("viewer")).status_code == 403
        for name, content in files:
            assert api.post(base, files={"file": (name, content)}, headers=auth("analyst")).status_code == 201
        assert api.post(base, files={"file": (files[0][0], files[0][1])}, headers=auth("analyst")).status_code == 409           # the same file twice
        assert api.post(base, files={"file": ("x.xlsx", b"not a workbook")}, headers=auth("analyst")).status_code == 422
        items = api.get(base, headers=auth("viewer")).json()
        assert items[0]["id"] == "all" and any(i["id"] == item for i in items)
        rep = api.get(f"{base}/{item}/report?lang=en", headers=auth("viewer")).json()
        assert rep["report"]["sections"][0]["key"] == "summary" and "analysis" in rep
        assert api.get(f"{base}/{item}/report.pdf?lang=ar", headers=auth("viewer")).content[:4] == b"%PDF"
        assert api.get(f"{base}/{item}/report.xlsx?lang=en", headers=auth("viewer")).content[:2] == b"PK"
        assert api.delete(f"{base}/all", headers=auth("admin")).status_code == 409                                              # history is never deleted
        assert api.get(f"/api/settings/{key}.thresholds", headers=auth("viewer")).status_code == 200
        assert api.put(f"/api/settings/{key}.thresholds", json={"nope": 1}, headers=auth("admin")).status_code == 422
    assert api.put("/api/settings/vehicles.plates", json={"aliases": ["bad"]}, headers=auth("admin")).status_code == 422
    assert api.put("/api/settings/vehicles.plates", json={"aliases": ["ع ن 3333=ع م 3333"]}, headers=auth("analyst")).status_code == 403
    assert api.put("/api/settings/vehicles.plates", json={"aliases": ["ع ن 3333=ع م 3333"]}, headers=auth("admin")).status_code == 200
    assert api.get("/api/settings/vehicles.plates", headers=auth("viewer")).json()["effective"]["aliases"] == ["ع ن 3333=ع م 3333"]


def test_rent_months_the_file_has_no_column_for_are_not_available_not_zero():
    base = {"governorate": "g", "scope": "branch", "start": None, "end": None, "start_raw": None, "end_raw": None, "advance": None, "deposit": None, "contract_rent": None,
            "current_rent_stated": None, "landlord": None, "flags": [], "versions": 1}
    contracts = [{**base, "key": "a", "name": "a", "months": {"2025-01": 100.0, "2025-04": 500.0}, "nopay": set()}, {**base, "key": "b", "name": "b", "months": {"2025-01": 50.0}, "nopay": {"2025-04"}}]
    a = rent_engine.analyze(contracts, effective_thresholds.__globals__["default_thresholds"]("rent"))
    assert [m["available"] for m in a["months"]] == [True, False, False, True]
    assert a["totals"]["missing"] == ["2025-02", "2025-03"] and a["totals"]["year_months"] == 2          # only the months that exist are counted
    assert a["months"][3]["gap_before"] is True and a["months"][3]["large"] is False                     # no «large change» claim across a gap
    assert a["totals"]["yoy_pct"] is None


def test_rent_month_with_entries_for_only_a_few_contracts_is_partial_and_not_compared():
    base = {"governorate": "g", "scope": "branch", "start": None, "end": None, "start_raw": None, "end_raw": None, "advance": None, "deposit": None, "contract_rent": None,
            "current_rent_stated": None, "landlord": None, "flags": [], "versions": 1, "nopay": set()}
    contracts = [{**base, "key": str(i), "name": f"c{i}", "months": {"2025-01": 100.0, "2025-03": 110.0, **({"2025-02": 100.0} if i == 0 else {})}} for i in range(10)]
    a = rent_engine.analyze(contracts, effective_thresholds.__globals__["default_thresholds"]("rent"))
    assert [(m["period"], m.get("partial")) for m in a["months"]] == [("2025-01", False), ("2025-02", True), ("2025-03", False)]
    assert a["totals"]["partial_months"] == ["2025-02"] and round(a["totals"]["dpct"], 1) == 10.0        # March is compared with January (the last complete month), not with the partial February
    assert a["totals"]["year_months"] == 2
    from app.modules.rent_analysis import report as rent_report
    th = effective_thresholds.__globals__["default_thresholds"]("rent")
    rm = rent_report.build_report(a, {"versions": [{"file": "f.xlsx", "id": 1, "layout": "rent_register", "uploaded_at": "", "from": None, "to": None, "records": 0, "added": None, "dropped": None, "changed": 0}],
                                      "changes": [], "summaries": [{}]}, th, {k: "default" for k in th}, "en", {}, [], False, None)
    assert any("partial" in i["text"].lower() for i in rm.sections[0].insights)


def test_rent_conflicting_copies_are_resolved_by_evidence_or_excluded_and_listed_with_their_sources(session):
    from tests.ops_helpers import rent_exceptions_workbook
    p = rent_wb.parse(rent_exceptions_workbook())
    a = rent_wb.assemble(p)
    names = [c["copy"].name for c in a["contracts"]]
    assert names == ["منشأة س"]                                                                       # the undated rows were folded into the dated, formula-owned record: one contract, not three
    assert a["contracts"][0]["governorate"] == "بيتا" and a["contracts"][0]["evidence"] == "sheet_formula"
    assert [e["name"] for e in a["excluded"]] == ["عقد متعارض"] and issue(p, "copies_conflict_unowned")  # differing copies, no owner: not valued
    ex = {e["name"]: e for e in a["exceptions"]}
    assert ex["منشأة س"]["status"] == "resolved" and "undated_copies_merged" in ex["منشأة س"]["evidence"]
    assert ex["عقد متعارض"]["status"] == "excluded"
    ad = RentAdapter()
    ad.ingest(session, rent_exceptions_workbook(), "exc.xlsx", "bob", {})
    d = rent_service.load(session)
    assert [c["name"] for c in d["contracts"]] == ["منشأة س"]                                          # excluded: not in the figures
    groups = next(e for e in d["summaries"][0]["exceptions"] if e["name"] == "عقد متعارض")["groups"]
    assert sorted(g["months"]["2025-01"] for g in groups) == [100.0, 200.0] and all(g["sheets"] for g in groups)   # the conflicting values and their sources are kept
    rm, _a = ad.build(session, "all", "en", True, {})
    q = next(s for s in rm.sections if s.key == "quality")
    tbl = next(t for t in q.tables if t["key"] == "exceptions")
    assert any("Excluded" in r["st"] for r in tbl["rows"]) and any("200" in r["v"] for r in tbl["rows"])


def test_overtime_excluded_sheet_is_listed_and_a_later_correct_month_replaces_not_available(session):
    from tests.ops_helpers import overtime_with_march
    ad = OvertimeAdapter()
    ad.ingest(session, overtime_workbook(), "ot.xlsx", "bob", {})
    d1 = ot_service.load(session)
    assert d1["summaries"][0]["excluded_sheets"] == [{"sheet": "مارس", "period": "2025-02", "reason": "duplicate_period_sheet", "same_as": "فبراير", "employees": 3}]
    a1 = ot_engine.analyze(d1["rows"], d1["roster"], effective_thresholds.__globals__["default_thresholds"]("overtime"))
    assert {r["period"] for r in d1["rows"]} == {"2025-01", "2025-02"} and a1["totals"]["all_hours"] == 107     # the repeated sheet adds nothing to any total
    rm, _a = ad.build(session, "all", "en", True, {})
    q = next(s for s in rm.sections if s.key == "quality")
    assert any(t["key"] == "excluded" and t["rows"][0]["s"] == "مارس" for t in q.tables)
    # the correct March arrives later: a new version; the earlier versions are kept and the month becomes available
    n_before = session.scalar(select(func.count()).select_from(OpRecord))
    ad.ingest(session, overtime_with_march(), "ot_march.xlsx", "bob", {})
    d2 = ot_service.load(session)
    assert {r["period"] for r in d2["rows"]} == {"2025-01", "2025-02", "2025-03"}
    assert session.scalar(select(func.count()).select_from(OpRecord)) > n_before and [v["layout"] for v in d2["versions"]] == ["overtime_monthly"] * 2
    a2 = ot_engine.analyze(d2["rows"], d2["roster"], effective_thresholds.__globals__["default_thresholds"]("overtime"))
    assert a2["totals"]["missing"] == [] and [m["hours"] for m in a2["avail"]][-1] == 12 and d2["changes"] == []     # nothing earlier changed; March simply became available


# ================================================================================================ annual report
@pytest.fixture
def three_modules(session):
    RentAdapter().ingest(session, rent_workbook(1, 2026), "rent26.xlsx", "bob", {})
    VehicleAdapter().ingest(session, repairs_workbook(), "repairs.xlsx", "bob", {})
    OvertimeAdapter().ingest(session, overtime_workbook(), "ot.xlsx", "bob", {})
    return session


def test_annual_report_links_rent_and_fleet_in_money_and_keeps_overtime_in_hours(three_modules):
    from app.modules.annual_report.adapter import AnnualAdapter
    from app.modules.annual_report import engine as annual_engine
    s = three_modules
    ad = AnnualAdapter()
    assert [i["id"] for i in ad.list_items(s)] == ["y:2026", "y:2025"]
    rm, A = ad.build(s, "y:2026", "en", True, {})
    cb = A["combined"]
    assert cb["common"] == ["2026-01", "2026-02"]                                                    # both rent and fleet complete only in Jan–Feb (the fleet has no March)
    assert [round(x["combined"]) for x in cb["rows"] if x["combined"] is not None] == [4880 + 8120, 2680 + 4310]
    assert round(cb["total"]) == 4880 + 8120 + 2680 + 4310 and round(cb["rent"]) == 7560 and round(cb["fleet"]) == 12430
    mar = next(x for x in cb["rows"] if x["period"] == "2026-03")
    assert mar["rent_state"] == "ok" and mar["fleet_state"] == "missing" and mar["combined"] is None    # no combined figure for a month one component lacks
    assert A["modules"]["overtime"] is None                                                           # the overtime file is a 2025 statement
    keys = [x.key for x in rm.sections]
    assert keys == ["summary", "cost", "rent", "fleet", "vehicles", "overtime", "compare", "drivers", "moves", "signals", "kpis", "quality", "versions"]
    text = " ".join(i["text"] for sec in rm.sections for i in sec.insights)
    assert "Overtime: no data for this year" in text
    # same numbers as the module dashboards (same source, same engine)
    rent_a = RentAdapter().build(s, "all", "en", True, {})[1]
    veh_a = VehicleAdapter().build(s, "all", "en", True, {})[1]
    rm_by = {m["period"]: m["total"] for m in rent_a["months"] if m.get("available")}
    fm_by = {m["period"]: m["total"] for m in veh_a["months"]}
    for x in cb["rows"]:
        if x["combined"] is not None:
            assert x["rent"] == rm_by[x["period"]] and x["fleet"] == fm_by[x["period"]]
    assert A["modules"]["fleet"]["total"] == veh_a["totals"]["total"] and A["modules"]["rent"]["total"] == sum(m["total"] for m in rent_a["avail"])
    # the dashboard API of the annual module reports the same figures
    assert ad.api_analysis(A)["combined"]["total"] == cb["total"]
    assert PdfExporter().render(rm)[:4] == b"%PDF" and ExcelExporter().render(rm)[:2] == b"PK"


def test_annual_report_2025_has_overtime_hours_no_money_total_and_lists_exclusions(three_modules):
    from app.modules.annual_report.adapter import AnnualAdapter
    s = three_modules
    rm, A = AnnualAdapter().build(s, "y:2025", "en", False, {})
    assert A["combined"]["common"] == [] and A["combined"]["total"] == 0                               # overtime is never added; there is no fleet data for 2025
    ot = A["modules"]["overtime"]
    assert round(ot["hours"], 1) == 107.0 and [m["period"] for m in ot["months"]] == ["2025-01", "2025-02"]
    excl = next(t for sec in rm.sections if sec.key == "quality" for t in sec.tables if t["key"] == "exclusions")["rows"]
    assert any("Overtime" in r["m"] and "excluded" in r["i"].lower() for r in excl)                    # the repeated March sheet
    emp = next(t for sec in rm.sections if sec.key == "overtime" for t in sec.tables if t["key"] == "ot_emp")
    assert "name" not in [c["key"] for c in emp["columns"]]                                           # names are admin-only
    rm_admin, _ = AnnualAdapter().build(s, "y:2025", "en", True, {})
    emp_a = next(t for sec in rm_admin.sections if sec.key == "overtime" for t in sec.tables if t["key"] == "ot_emp")
    assert "name" in [c["key"] for c in emp_a["columns"]]
    vers = next(t for sec in rm.sections if sec.key == "versions" for t in sec.tables if t["key"] == "versions")["rows"]
    assert {r["m"] for r in vers} == {"Rent", "Fleet", "Overtime"}


def test_annual_api(api, three_modules):
    base = "/api/analysis/annual/datasets"
    assert api.post(base, files={"file": ("x.xlsx", b"x")}, headers=auth("analyst")).status_code == 422
    assert [i["id"] for i in api.get(base, headers=auth("viewer")).json()] == ["y:2026", "y:2025"]
    rep = api.get(f"{base}/y:2026/report?lang=ar", headers=auth("viewer")).json()
    assert rep["analysis"]["combined"]["common"] == ["2026-01", "2026-02"] and rep["report"]["sections"][0]["key"] == "summary"
    assert api.get(f"{base}/y:2026/report.pdf?lang=ar", headers=auth("viewer")).content[:4] == b"%PDF"
    assert api.get(f"{base}/y:2026/report.xlsx?lang=en", headers=auth("viewer")).content[:2] == b"PK"
    assert api.get(f"{base}/y:1999/report", headers=auth("viewer")).status_code == 404
    assert api.delete(f"{base}/y:2026", headers=auth("admin")).status_code == 409
    assert api.get("/api/settings/annual.thresholds", headers=auth("viewer")).status_code == 200


def test_plate_aliases_are_a_setting_not_code_and_can_be_changed_later(session):
    from app.core.settings_store import set_setting
    ad = VehicleAdapter()
    ad.ingest(session, repairs_workbook(False, "ع م 3333"), "repairs.xlsx", "bob", {})       # the repairs statement writes the plate with a different letter
    ad.ingest(session, usage_book(), "usage.xlsx", "bob", {})                                  # «ع ن 3333»
    ad.ingest(session, card_book(), "card.xlsx", "bob", {})                                    # «3333» alone
    d = veh_service.load(session)
    assert sorted(v["plate"] for v in d["vehicles"]) == sorted(["س ص 2222", "ع م 3333", "ع ن 3333", "3333"]) and d["candidates"][0]["number"] == "3333"   # nothing linked by guessing
    set_setting(session, "vehicles.plates", {"aliases": ["ع ن 3333=ع م 3333"]}, "alice")      # approved by the owner in the settings
    d = veh_service.load(session)
    one = [v for v in d["vehicles"] if v["key"] == plates.key("ع م 3333")]
    assert len(one) == 1 and sorted(one[0]["cost"]) == ["2026-01", "2026-02"] and sorted(one[0]["usage"]) == ["2026-01", "2026-02"] and sorted(one[0]["service"]) == ["2026-01", "2026-02"]
    assert not d["candidates"] and len(d["vehicles"]) == 2
    set_setting(session, "vehicles.plates", {"aliases": []}, "alice")                          # and it can be undone: the stored records never changed
    assert len(veh_service.load(session)["vehicles"]) == 4


def test_production_refuses_to_start_without_an_admin_token(monkeypatch, engine):
    import pytest as _pytest
    from fastapi.testclient import TestClient
    from app.config import get_settings
    import app.main as main_mod
    from app.main import app
    monkeypatch.setattr(main_mod, "get_engine", lambda: engine)
    for env, tokens, ok in (("production", "", False), ("production", "a:viewer:v", False), ("production", "a:admin:x", True), ("development", "", True)):
        monkeypatch.setenv("APP_ENV", env)
        monkeypatch.setenv("API_TOKENS", tokens)
        get_settings.cache_clear()
        if ok:
            with TestClient(app):
                pass
        else:
            with _pytest.raises(RuntimeError):
                with TestClient(app):
                    pass
    monkeypatch.delenv("APP_ENV")
    monkeypatch.delenv("API_TOKENS")
    get_settings.cache_clear()


def test_vehicle_cost_cells_cleared_in_a_newer_version_are_not_carried_over(session):
    """A newer file that moves a cost from one category to another must not leave the old category's value in the current view (it would double-count)."""
    ad = VehicleAdapter()
    ad.ingest(session, repairs_workbook(), "repairs.xlsx", "bob", {})
    wb = repairs_workbook(True)                                    # Feb: the first vehicle's repair 700 -> 900 (same category) …
    from openpyxl import load_workbook
    import io
    book = load_workbook(io.BytesIO(wb))
    ws = book["فبراير"]
    ws["E5"].value, ws["F5"].value = None, 700                     # … and here the 700 is moved from «صيانة» to «اخري - تكييف» instead
    ws["H5"].value = 710
    buf = io.BytesIO()
    book.save(buf)
    ad.ingest(session, buf.getvalue(), "repairs_moved.xlsx", "bob", {})
    d = veh_service.load(session)
    v = next(x for x in d["vehicles"] if x["plate"] == "س ص 1111")
    cats = {k: val for k, val in v["cost"]["2026-02"]["values"].items() if k.startswith("cat:")}
    assert cats.get("cat:اخري - تكييف") == 700 and "cat:صيانة" not in cats and sum(cats.values()) == 710       # moved, not duplicated
    gone = [c for c in d["changes"] if c["field"] == "cat:صيانة" and c["period"] == "2026-02"]
    assert [(c["old"], c["new"]) for c in gone] == [(700, None)]                                            # and the move is visible as a change
