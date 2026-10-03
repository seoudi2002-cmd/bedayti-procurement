"""Synthetic workbooks that mimic the structure of the custody report files (no real data)."""
import io

from openpyxl import Workbook

DESC = "اغلاق عهدة مؤقتة طرف شخص تجريبي {n} لسداد مصروفات"


def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def gl_workbook(lines, with_year_title=True, pivot_total_wrong=False) -> bytes:
    """lines: (amount, account_no, class1, cost_center, branch, description, je, month)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Amount", "Acc.name", "Class 1", "class 2", "branch", "deccription", "number", "month"])
    for ln in lines:
        ws.append(list(ln))
    months = sorted({ln[7] for ln in lines})
    cats = sorted({ln[2] for ln in lines})
    p2 = wb.create_sheet("Sheet2")
    p2["A3"] = "Sum of Amount"
    p2["A4"] = "Row Labels"
    for j, m in enumerate(months, 2):
        p2.cell(4, j, m)
    p2.cell(4, len(months) + 3, "Grand Total")
    for i, c in enumerate(cats, 5):
        p2.cell(i, 1, c)
        for j, m in enumerate(months, 2):
            v = sum(ln[0] for ln in lines if ln[2] == c and ln[7] == m)
            if v:
                p2.cell(i, j, v)
    gt = 5 + len(cats)
    p2.cell(gt, 1, "Grand Total")
    for j, m in enumerate(months, 2):
        p2.cell(gt, j, sum(ln[0] for ln in lines if ln[7] == m) + (1000 if pivot_total_wrong and j == 2 else 0))
    p2.cell(gt, len(months) + 3, sum(ln[0] for ln in lines))
    p3 = wb.create_sheet("Sheet3")
    p3["A1"] = "Temperory Custodies Analysis From Jan up To Mar-2026" if with_year_title else "Temporary custody analysis"
    p3["A2"], p3["B2"] = "نوع المصروف", "Month"
    for i, c in enumerate(cats, 3):
        p3.cell(i, 1, "علاقات عامة" if c.startswith("Public") else "ضيافة" if c.startswith("Cat") else f"بند {i}")
        p3.cell(i, 2, c)
    return _save(wb)


def standard_gl() -> bytes:
    rows = [
        (100000, 51080000, "Governmental Fees", 3401, "Head Office ", DESC.format(n=1), "JE301", 1),
        (40000, 51030000, "Catring & Cleaning Exp", 3401, "Head Office ", DESC.format(n=1), "JE301", 1),
        (20000, 51200000, "Transportation & Carriage Exp", 2101, "Branch One", DESC.format(n=2), "JE302", 1),
        (15000, 51200000, "Transportation & Carriage Exp", 2102, "Branch Two", DESC.format(n=3), "JE303", 1),
        (60000, 51080000, "Governmental Fees", 3401, "Head Office", DESC.format(n=4), "JE301", 2),
        (500000, 51080000, "Governmental Fees", 2101, "Branch One", DESC.format(n=2), "JE302", 2),
        (12000, 51030000, "Catring & Cleaning Exp", 2102, "Branch Two", DESC.format(n=3), "JE303", 2),
        (10000, 51030000, "Catring & Cleaning Exp", 2103, "Branch Three", DESC.format(n=5), "JE304", 2),
        (10000, 51030000, "Catring & Cleaning Exp", 2103, "Branch Three", DESC.format(n=5), "JE304", 2),
        (80000, 51370000, "Public Relation", 3401, "Head Office", DESC.format(n=4), "JE305", 3),
        (9000, 51370000, "Public Relations", 2102, "Branch Two", DESC.format(n=3), "JE306", 3),
        (-2500, 51080000, "Governmental Fees", 3401, "Head Office", "Reclase Entry Between (Governmental&Utilites)Exp", "GL", 3),
        (7000, 51200000, "Transportation & Carriage Exp", None, "Branch One", DESC.format(n=2), "JE302", 3),
        (3000, 51200000, "Transportation & Carriage Exp", 2101, "Branch Uno", DESC.format(n=2), "JE302", 3),
        (8000, 51200000, "Transportation & Carriage Exp", 2301, "فروع قنا", DESC.format(n=6), "JE307", 3),
    ]
    return gl_workbook(rows)


def branch_wide_sheet(wb, title, month_text, rows, total_wrong_row=None, stray=False):
    ws = wb.create_sheet(title)
    ws["A1"], ws["B1"], ws["C1"] = "", "البيان", "تاريخ الصرف"
    ws["D1"], ws["E1"], ws["G1"], ws["J1"] = "TRANSPORTATION & CARRIAGE", "CATRING & CLEANING", "UTILITIES", "الإجمالي"
    ws["D2"], ws["E2"], ws["F2"], ws["G2"], ws["H2"] = "انتقالات", "نظافة", "ضيافه", "كهرباء", "مياه"
    ws.merge_cells("E1:F1")
    ws.merge_cells("G1:H1")
    ws["J1"] = None
    ws["I1"] = "الإجمالي"
    r = 3
    first = r
    for i, (name, vals) in enumerate(rows, 1):
        ws.cell(r, 1, i)
        ws.cell(r, 2, name)
        ws.cell(r, 3, month_text)
        for c, v in zip(range(4, 9), vals):
            ws.cell(r, c, v)
        total = sum(v or 0 for v in vals)
        ws.cell(r, 9, total + 100 if total_wrong_row == i else total)
        r += 1
    ws.cell(r, 1, "الإجمالي الفرعي")
    for c in range(4, 10):
        ws.cell(r, c, sum((ws.cell(k, c).value or 0) for k in range(first, r)))
    if stray:
        ws.cell(first + 1, 12, 123456)
    return ws


def monthly_branch_workbook() -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    names = [f"فرع {i}" for i in range(1, 13)]
    for title, mtxt, mult in (("شهر يناير", "يناير 2026", 1), ("شهر مارس", "مارس 2026", 3)):
        rows = [(n, [100 * mult * (i + 1), None, 50 * mult, 25, 10]) for i, n in enumerate(names)]
        branch_wide_sheet(wb, title, mtxt, rows, total_wrong_row=2 if mult == 1 else None, stray=mult == 3)
    # February: English pivot layout with other branch names (like a re-typed sheet)
    ws = wb.create_sheet("شهر فبراير")
    ws.append(["BRANCH", "Catring & Cleaning Exp", "Utilities Exp", " Total"])
    for i in range(1, 13):
        ws.append([f"Branch {i}", 200 * i, 30, 200 * i + 30])
    tot = wb.create_sheet("total")
    tot.append(["الشهر", "Total"])
    tot.append(["يناير", 99999])  # deliberately different from the computed total
    return _save(wb)


def monthly_custodian_workbook() -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("تحليلي المركز الرئيسي شهر يناير")  # layout A: rows = expense accounts, columns = custodians
    ws.append(["ACC.NAME", "المركز الرئيسي - شخص أ", "المركز الرئيسي - شخص ب", "الاجمالي", "نسبة المصروف للرئيسي"])
    ws.append(["Meals Exp", 1000, 500, 1500, 0.6])
    ws.append(["Fuel & Vehicles Exp", None, 1000, 1000, 0.4])
    ws.append(["الاجمالي", 1000, 1500, 2500])
    ws2 = wb.create_sheet("تحليلي المركزالرئيسي شهر فبراير")  # layout B: rows = custodians, columns = categories
    ws2.append(["Row Labels", "Meals Exp", "Fuel & Vehicles Exp", " Total"])
    ws2.append(["(  شخص أ  )", 700, None, 700])
    ws2.append(["(  شخص ب  )", 300, 2000, 2300])
    ws2.append(["Total", 1000, 2000, 3000])
    return _save(wb)
