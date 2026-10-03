"""Synthetic reference files (no real data): an asset-register export with the same header as the real one, and a native-text PDF."""
import functools
import io
from datetime import datetime

from openpyxl import Workbook

HEADER = ["Asset Book", "Asset Number", "Description", "Tag Number", "Serial Number", "Location", "Major Category", "Category Segments", "Accounting Date", "Date Placed In Service",
          "Prorate Date", "Prorate Convention Code", "Deprn Start Date", "Date Retired", "Asset Type", "Method Code", "Life In Months", "Current Units", "Current Period",
          "Original Cost", "Adjusted Cost", "Recoverable Cost", "Cost", "Ytd Accu Depreciation", "Tal Accu Depreciation Last", "D Net Book Value Last",
          *[f"{m} System Deprn Amount" for m in ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")]]


def _row(no, desc, loc, cat, seg, cost, accu, in_service, acct=datetime(2024, 6, 23), period=datetime(2024, 9, 24), units=1, serial=None, ytd=None, months=None):
    months = months if months is not None else [0, 7, 7, 7, 7, 7, 7, 7, 7, 0, 0, 0]
    return ["BOOK", no, desc, None, serial, loc, cat, seg, acct, in_service, datetime(2022, 1, 3), "CONV-DATE", datetime(2022, 1, 3), None, "Capitalized", "STL", "\xa0 010.00", units,
            period, cost, cost - 1, cost - 1, cost, ytd if ytd is not None else sum(months), accu, cost - accu, *months]


@functools.lru_cache(maxsize=None)  # saved workbooks embed a timestamp: identical bytes are needed for duplicate-upload tests
def asset_workbook(version: int = 1) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    ws.append(_row("10000", "Chair", "Gov1-CityA-CityA Office-", "Chairs-1", "Furniture", 800, 200, datetime(2022, 1, 2)))
    ws.append(_row("10001", "Chair", "Gov1-CityA-CityA Office-", "Chairs-1", "Furniture", 800, 200, datetime(2022, 1, 2)))
    ws.append(_row("10002", "PC", "Cairo-Cairo-Head Office-", "PC-1", "Computers & Software", 20000, 9000, "31-07-2024", acct=datetime(2024, 6, 23)))     # in-service as TEXT, after accounting date
    ws.append(_row("10003", "PC", "Cairo-Cairo-Head Office-", "PC-1", "Computers & Software", 20000, 9000, datetime(2022, 1, 2), acct=datetime(2024, 12, 23)))   # accounting after period
    ws.append(_row("10004", "Table", "OddLocationText", "Table-1", "Furniture", 0, 0, datetime(2022, 1, 2)))                                                   # zero cost, unsplit location
    ws.append(_row("10005", "Laptop", "Gov2-CityB-CityB Office-", "Laptop-1", "Computers & Software", 30000, 12000, datetime(2023, 1, 5), serial="SN1"))
    ws.append(_row("10005", "Laptop", "Gov2-CityC-CityC Office-", "Laptop-1", "Computers & Software", 30000, 12000, datetime(2023, 1, 5)))                       # duplicate number
    ws.append(_row("10006", "Broken", "Gov2-CityB-CityB Office-", "Desks-1", "Furniture", 1000, 100, datetime(2022, 1, 2), ytd=999))                           # Ytd != sum of months
    if version == 2:        # a later export: one asset gone, one new, one moved and re-costed
        ws.delete_rows(2)
        ws.append(_row("10007", "New server", "Cairo-Cairo-Head Office-", "Server-1", "Computers & Software", 90000, 1000, datetime(2024, 8, 1), period=datetime(2024, 12, 31)))
        ws.cell(4, 6).value = "Gov3-CityD-CityD Office-"
        ws.cell(4, 23).value = 21000
        for r in range(2, ws.max_row + 1):          # a later export: every row carries the later period
            ws.cell(r, 19).value = datetime(2024, 12, 31)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@functools.lru_cache(maxsize=None)
def regulation_pdf(pages: int = 3) -> bytes:
    """A native-text PDF (the extraction job reads the text layer; no OCR needed)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for i in range(1, pages + 1):
        c.setFont("Helvetica", 12)
        c.drawString(60, 780, f"Regulation sample - Article {i}")
        c.drawString(60, 760, "The branch manager shall keep a register of the assets of the branch.")
        c.showPage()
    c.save()
    return buf.getvalue()
