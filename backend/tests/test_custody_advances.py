from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.core.analysis import AnalysisError
from app.core.reporting.excel import ExcelExporter
from app.core.reporting.pdf import PdfExporter
from app.core.settings_store import set_setting
from app.models import CustodyAdvance
from app.modules.custody_analysis import advances, layouts
from app.modules.custody_analysis.adapter import CustodyAdapter
from tests.custody_helpers import advance_register_workbook, standard_gl
from tests.test_custody_analysis import api, auth  # noqa: F401  (fixtures)


def _parse(**kw):
    import io
    return advances.parse_advances(load_workbook(io.BytesIO(advance_register_workbook(**kw)), data_only=True))


def test_reader_states_flags_controls_and_never_guesses():
    p = _parse()
    assert len(p.rows) == 6 and [c.label for c in p.controls] == ["اقفال شهر يناير", "اقفال شهر فبراير"]      # closing rows are controls, not advances
    st = {x.holder: x.state for x in p.rows}
    assert st == {"Holder A": "settled", "Holder B": "settled", "Holder C": "open", "Holder D": "settled_date_unreadable", "Holder E": "refunded_note", "Holder F": "settled"}
    assert next(x for x in p.rows if x.holder == "Holder D").settled_on is None                      # an unreadable date stays unknown
    codes = {i.code for i in p.issues}
    assert {"settlement_date_invalid", "settled_before_advance", "closed_by_status_note", "control_total_differs", "closing_second_figure_unexplained"} <= codes
    assert advances.as_of_of(p.rows) == date(2026, 4, 16)
    assert "control_total_differs" not in {i.code for i in _parse(break_closing=False).issues}      # an agreeing closing row raises nothing
    with pytest.raises(layouts.UnrecognisedLayout):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.append(["a", "b"])
        advances.parse_advances(wb)


def _upload(session):
    ad = CustodyAdapter()
    r = ad.ingest(session, advance_register_workbook(), "advances.xlsx", "bob", {})
    return ad, r


def test_ingest_detects_the_register_and_rejects_repeats(session):
    ad, r = _upload(session)
    assert r.meta["layout"] == "advance_register" and r.meta["facts"] == 6
    assert session.query(CustodyAdvance).count() == 6
    with pytest.raises(AnalysisError) as e:
        ad.ingest(session, advance_register_workbook(), "again.xlsx", "bob", {})
    assert e.value.status == 409
    assert ad.list_items(session)[0]["layout"] == "advance_register"
    ad2 = ad.ingest(session, standard_gl(), "gl.xlsx", "bob", {})                              # the other layouts are unaffected
    assert ad2.meta["layout"] == "gl_settlement_lines"


def test_figures_ageing_lag_and_exceptions(session):
    ad, r = _upload(session)
    set_setting(session, "custody_advances.thresholds", {"repeat_advances": 2}, "t")
    rm, a = ad.build(session, r.item_id, "en", True, {})
    t = a["totals"]
    assert (t["n"], t["issued"], t["closed_n"], t["open_n"], t["open_amount"]) == (6, Decimal(315000), 5, 1, Decimal(50000))
    assert t["undated_closed_n"] == 2 and t["open_pct"] == pytest.approx(50000 / 315000 * 100)
    assert a["as_of"] == date(2026, 4, 16)
    assert a["lag"]["n"] == 2 and a["lag"]["max"] == 102 and a["lag"]["median"] == pytest.approx(53.5) and a["lag"]["long"] == 1      # the settled-before row has no lag
    assert [(w["key"], w["n"], w["amount"]) for w in a["ageing"]] == [("a1", 0, 0), ("a2", 0, 0), ("a3", 1, Decimal(50000)), ("a4", 0, 0)]    # open since 20 Jan: 86 days at 16 Apr
    jan, feb = a["months"][0], a["months"][1]
    assert (jan["issued_n"], jan["issued"], jan["settled"], jan["open"]) == (3, Decimal(80000), Decimal(20000), Decimal(60000))
    assert (feb["issued_n"], feb["issued"], feb["settled"]) == (3, Decimal(235000), Decimal(200000))
    assert [m["period"] for m in a["months"]] == [(2026, 1), (2026, 2), (2026, 4)] and a["months"][2]["settled"] == Decimal(10000)
    assert {k: len(v) for k, v in a["exceptions"].items() if v} == {"large": 1, "long_settlement": 1, "settled_before": 1, "closed_undated": 2}
    assert [b["label"] for b in a["repeating"]] == ["فرع ألف"]
    byk = {b["label"]: b for b in a["branches"]}
    assert byk["فروع الغربية"]["kind"] == "group" and byk["المركز الرئيسي"]["kind"] == "head_office" and byk["فروع الغربية"]["open"] == Decimal(50000)
    c = {k["label"]: k for k in a["controls"]}
    assert c["اقفال شهر يناير"]["stated"] == c["اقفال شهر يناير"]["computed"] == Decimal(80000) and c["اقفال شهر يناير"]["other_vs_balance"] == 0   # 2nd figure = dated month-end balance
    assert c["اقفال شهر فبراير"]["stated"] == Decimal(999) and c["اقفال شهر فبراير"]["computed"] == Decimal(235000)
    assert [s.key for s in rm.sections] == ["summary", "monthly", "branches", "ageing", "quality"]
    assert any(i["metric"] == "second_figure" for i in rm.meta["summary_items"])


def test_personal_names_only_for_admins_and_filters_exports(session):
    ad, r = _upload(session)
    admin_rm, _ = ad.build(session, r.item_id, "en", True, {})
    viewer_rm, _ = ad.build(session, r.item_id, "en", False, {})
    assert "Holder C" in str(admin_rm) and "Holder" not in str(viewer_rm).replace("Holder (personal data)", "")
    assert "Holder (personal data)" not in str(viewer_rm)
    dims = {d["key"]: d for d in viewer_rm.meta["filters"]["dimensions"]}
    assert set(dims) == {"periods", "branches", "states"} and [i["id"] for i in dims["periods"]["items"]] == ["2026-01", "2026-02"]
    _rm, a = ad.build(session, r.item_id, "en", True, {"states": ["open"]})
    assert a["totals"]["n"] == 1 and a["filtered"] and a["as_of"] == date(2026, 4, 16)
    _rm, a = ad.build(session, r.item_id, "en", True, {"periods": ["2026-02"]})
    assert a["totals"]["n"] == 3 and a["totals"]["issued"] == Decimal(235000)
    with pytest.raises(AnalysisError):
        ad.build(session, r.item_id, "en", True, {"branches": ["name:nope"]})
    for lang in ("ar", "en"):
        rm, _a = ad.build(session, r.item_id, lang, True, {})
        assert PdfExporter().render(rm)[:4] == b"%PDF" and ExcelExporter().render(rm)[:2] == b"PK"


def test_api_upload_thresholds_and_delete(api):
    base = "/api/analysis/custody/datasets"
    assert api.post(base, files={"file": ("a.xlsx", advance_register_workbook())}, headers=auth("viewer")).status_code == 403
    r = api.post(base, files={"file": ("a.xlsx", advance_register_workbook())}, headers=auth("analyst"))
    assert r.status_code == 201, r.text
    ds_id = r.json()["id"]
    rep = api.get(f"{base}/{ds_id}/report?lang=en&state=open", headers=auth("viewer")).json()
    assert rep["analysis"]["totals"]["n"] == 1 and "Holder" not in str(rep["report"])
    assert api.get(f"{base}/{ds_id}/report.pdf?lang=ar", headers=auth("viewer")).content[:4] == b"%PDF"
    assert api.put("/api/settings/custody_advances.thresholds", json={"top_n": 5}, headers=auth("analyst")).status_code == 403
    assert api.put("/api/settings/custody_advances.thresholds", json={"top_n": 5}, headers=auth("admin")).status_code == 200
    assert api.put("/api/settings/custody_advances.thresholds", json={"nope": 1}, headers=auth("admin")).status_code == 422
    assert api.delete(f"{base}/{ds_id}", headers=auth("admin")).status_code in (200, 204)
    assert api.get(f"{base}/{ds_id}/report", headers=auth("viewer")).status_code == 404
