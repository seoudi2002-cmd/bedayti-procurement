"""Synthetic workbooks that mimic the *shape* of the real reference files (merged cells, subtotal rows, title
rows, error cells) without any real data."""
import io
from datetime import datetime

from openpyxl import Workbook


def xlsx(sheets: dict[str, list[list]], merges: dict[str, list[str]] | None = None, hidden: dict[str, list[str]] | None = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
        for rng in (merges or {}).get(name, []):
            ws.merge_cells(rng)
        for col in (hidden or {}).get(name, []):
            ws.column_dimensions[col].hidden = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


BRANCH_HEADER = ["م", "المنطقة", "مدير  المنطقة ", "اسم الفرع ", "العنــوان", "اسم مدير الفرع", "رقم تليفون الفرع",
                 "هاتف مدير الفرع ", "هاتف مدير الفرع ", "هاتف مدير المنطقة ", "البريد الإلكتروني للفرع"]


def branch_workbook() -> bytes:
    rows = [["بيـــانات الفـــروع"], BRANCH_HEADER,
            [1, "الشرقية", "مدير شرقية", "ابوحماد", "عنوان 1", "مدير فرع واحد", 553413434, 1092841323, None, None, "a@x.test"],
            [2, None, None, "كفر صقر", "عنوان 2", "مدير فرع اثنين", None, 1282138256, None, None, "b@x.test"],
            ["إجمالي منطقة الشرقية", None, None, 2],
            [3, "الغربية", None, "طنطا", None, None, None, None, None, None, "a@x.test"],
            ["إجمالي منطقة الغربية", None, None, 1]]
    return xlsx({"فروع بدايتي": rows, "Sheet1": [[None, None, "طنطا", "الغربية"]]},
                merges={"فروع بدايتي": ["B3:B4", "C3:C4"]}, hidden={"فروع بدايتي": ["G"]})


def employee_workbook(rows: list[list]) -> bytes:
    return xlsx({"Sheet3": [["HC", len(rows)], ["الكود", "الاسم", "تاريخ التعيين", "الوظيفة", "الفرع", "المحافظه"], *rows]})


SUPPLIER_HEADER = ["#", "اسم المورد", "رقم سجل\n  المورد", "كود المورد", "التصنيف", "الخدمات", "مسئول التواصل", "رقم الهاتف",
                   "البريد الالكتروني", "العنوان", "السجل التجاري", "البطاقة الضريبية", "ملاحظات"]

REQ_HEADER = ["م", "رقم اشعار الاحتياج", "تاريخ الطلب", "الادارة الطالبة", "وصف الاحتياج", "الاولوية", "حالة الاشعار", "ملاحظات"]
PO_HEADER = ["#", "رقم امر الشراء", "تاريخ\nامر الشراء", "اسم المورد", "رقم سجل\n  المورد", "تصنيف المورد", "تفاصيل الطلب",
             "رقم اشعار الاحتياج", "تصنيف امر الشراء", "Column 14", "اجمالي مبلغ الطلب", "حالة الامر", "المبلغ المسدد", "المتبقي"]
HANDOVER_HEADER = ["#", "نوع المدفوعة", "تاريخ الارسال للمالية", "رقم امر الشراء", "اسم المورد", "تصنيف المورد", "مضمون المذكرة",
                   "تصنيف الشراء", "اجمالي مبلغ الطلب"]


def d(y, m, dd):
    return datetime(y, m, dd)
