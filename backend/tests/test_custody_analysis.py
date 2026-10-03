import io
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import get_settings
from app.core.reporting.excel import ExcelExporter
from app.core.reporting.pdf import PdfExporter
from app.core.settings_store import set_setting
from app.db import get_session
from app.main import app
from app.modules.custody_analysis import layouts, service
from app.modules.custody_analysis.engine import F, analyze
from tests.custody_helpers import gl_workbook, monthly_branch_workbook, monthly_custodian_workbook, standard_gl

TH = {"mom_pct": 30, "mom_abs": 50000, "iqr_k": 1.5, "min_branches_for_iqr": 8, "line_share_of_category_pct": 40,
      "min_line_amount": 20000, "month_vs_avg_pct": 50, "min_months_for_trend": 4, "repeat_amount_min_count": 4,
      "repeat_amount_min_value": 1000, "top_n": 10, "dormant_new_min_amount": 5000}


# ------------------------------------------------------------------ parsing (F3 settlement journal)
def test_gl_parse_is_faithful_and_flags_without_fixing():
    p = layouts.parse_workbook(standard_gl())
    assert p.layout == "gl_settlement_lines" and p.scope_label == "temporary_custody"
    assert sum(f.amount for f in p.facts) == Decimal("871500.0000")  # every line, negative reclass included
    assert p.year == 2026 and p.year_source == "file"  # taken from the file's own title, not invented
    codes = {i.code for i in p.issues}
    assert {"negative_amount", "reclass_entry", "cost_center_missing", "cost_center_name_variants",
            "je_number_reused_across_months"} <= codes
    assert all(c.stated is not None and c.stated == c.computed for c in p.controls)  # file's pivots reconcile
    assert p.labels["Catring & Cleaning Exp"] == "ضيافة"  # source-provided display label
    ho = [f for f in p.facts if f.scope == "head_office"]
    assert ho and all(f.branch_kind == "head_office" for f in ho)
    grp = [f for f in p.facts if f.branch_kind == "group"]
    assert len(grp) == 1 and grp[0].scope == "branch"  # a branch group stays a group


def test_year_is_never_invented():
    p = layouts.parse_workbook(gl_workbook([(10, 1, "X", 1, "B", "d", "JE1", 1)], with_year_title=False))
    assert p.year is None and p.year_source == "none"
    p2 = layouts.parse_workbook(gl_workbook([(10, 1, "X", 1, "B", "d", "JE1", 1)], with_year_title=False), year=2027)
    assert p2.year == 2027 and p2.year_source == "uploader"


def test_pivot_control_mismatch_is_reported():
    p = layouts.parse_workbook(gl_workbook([(100, 1, "X", 1, "B", "d", "JE1", 1), (50, 1, "X", 1, "B", "d", "JE2", 2)], pivot_total_wrong=True))
    bad = [c for c in p.controls if c.stated != c.computed]
    assert len(bad) == 1 and bad[0].stated - bad[0].computed == Decimal("1000")


def test_unknown_workbook_is_refused_not_guessed():
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(["a", "b"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(layouts.UnrecognisedLayout):
        layouts.parse_workbook(buf.getvalue())


# ------------------------------------------------------------------ parsing (F2 / F1)
def test_branch_workbook_flags_disagreements_but_uses_components():
    p = layouts.parse_workbook(monthly_branch_workbook())
    assert p.layout == "monthly_branch_expense"
    codes = {i.code: i for i in p.issues}
    assert "row_total_differs_from_components" in codes  # row 2 January states +100
    assert "stray_cells_outside_table" in codes
    jan = sum(f.amount for f in p.facts if f.month == 1)
    assert jan == Decimal(sum(100 * (i + 1) + 50 + 25 + 10 for i in range(12)))  # components, not the wrong stated total
    wb_total = [c for c in p.controls if c.label.startswith("Month total")][0]
    assert wb_total.stated == Decimal("99999") and wb_total.computed != wb_total.stated
    # items keep the file's own two-level names: group (category) and sub-item
    item = next(f for f in p.facts if f.item == "كهرباء")
    assert item.category == "UTILITIES"


def test_custodian_workbook_both_layouts():
    p = layouts.parse_workbook(monthly_custodian_workbook())
    assert p.layout == "monthly_custodian_expense" and p.scope_label == "head_office"
    assert all(f.scope == "head_office" for f in p.facts)
    assert {c.month: c.computed for c in p.controls} == {1: Decimal("2500"), 2: Decimal("3000")}
    assert p.year is None  # no year anywhere in the file
    assert {f.holder for f in p.facts if f.month == 1} == {"المركز الرئيسي - شخص أ", "المركز الرئيسي - شخص ب"}


# ------------------------------------------------------------------ analysis engine
def _facts():
    def f(m, scope, br, cat, amt, ref):
        return F(2026, m, scope, br, br, "head_office" if scope == "head_office" else "branch", cat, None, None, Decimal(amt), ref)
    return [f(1, "head_office", "ho", "A", 1000, "r1"), f(1, "branch", "b1", "A", 3000, "r2"), f(1, "branch", "b2", "B", 1000, "r3"),
            f(2, "head_office", "ho", "A", 2000, "r4"), f(2, "branch", "b1", "B", 7000, "r5"), f(2, "branch", "b2", "B", 1000, "r6")]


def test_engine_totals_shares_and_variance():
    a = analyze(_facts(), TH)
    assert a["total"] == Decimal("15000")
    shares = {c["name"]: c["share"] for c in a["by_category"]}
    assert shares == {"A": pytest.approx(40.0), "B": pytest.approx(60.0)}
    assert sum(c["share"] for c in a["by_category"]) == pytest.approx(100.0)
    assert a["by_scope"]["head_office"]["total"] == Decimal("3000") and a["by_scope"]["branch"]["total"] == Decimal("12000")
    v = a["variance"]
    assert v["total"]["abs"] == Decimal("5000") and v["total"]["pct"] == pytest.approx(100.0)
    assert [r["name"] for r in a["by_branch"]] == ["b1", "b2"]
    assert a["by_branch"][0]["cum_share"] == pytest.approx(10000 / 12000 * 100)


def test_thresholds_are_inputs_not_constants():
    facts = _facts()
    loud = analyze(facts, {**TH, "mom_abs": 1000, "mom_pct": 10})
    quiet = analyze(facts, {**TH, "mom_abs": 10**9})
    assert any(o["rule"] == "mom_change" for o in loud["outliers"])
    assert not any(o["rule"] == "mom_change" for o in quiet["outliers"])


def test_unsupported_dimensions_are_listed_not_invented():
    facts = [F(2026, 1, "head_office", None, None, None, "A", None, "p", Decimal(5), "r")]
    a = analyze(facts, TH)
    assert a["by_scope"] is None and a["by_branch"] == []
    assert {"head_office_vs_branches", "branch_comparison", "top_branches", "branch_by_category"} <= set(a["unsupported"])


# ------------------------------------------------------------------ service + report
@pytest.fixture
def ds(session):
    return service.ingest(session, standard_gl(), "f3.xlsx", "tester")


def test_ingest_preserves_file_and_rejects_duplicates(session, ds):
    from pathlib import Path
    import hashlib
    assert hashlib.sha256(Path(ds.storage_path).read_bytes()).hexdigest() == ds.file_hash == hashlib.sha256(standard_gl()).hexdigest()
    with pytest.raises(service.DuplicateDataset):
        service.ingest(session, standard_gl(), "again.xlsx", "tester")
    facts = service.load_facts(session, ds.id)
    assert sum(f.amount for f in facts) == Decimal("871500")
    assert any(f.source_ref == "Sheet1!R2" for f in facts)  # lineage to the source row


@pytest.mark.parametrize("lang", ["ar", "en"])
def test_report_content_is_computed_and_honest(session, ds, lang):
    rm, a = service.build(session, ds, lang, admin=True)
    assert "871,500" in rm.executive_summary
    assert all(i["text"] for i in rm.meta["summary_items"])
    titles = [s.title for s in rm.sections]
    assert len(titles) >= 8
    keys = {t["key"] for s in rm.sections for t in s.tables}
    assert {"months", "categories", "category_month", "branches", "branch_category", "var_category", "outliers", "controls",
            "issues", "settings", "scope_month", "holders"} <= keys
    # thresholds are shown as defaults (not yet confirmed)
    st = next(t for s in rm.sections for t in s.tables if t["key"] == "settings")
    assert all(r["source"] == "default" for r in st["rows"])
    # the branch group is listed separately and not ranked as a branch
    assert not any("قنا" in r["name"] for r in next(t for s in rm.sections for t in s.tables if t["key"] == "branches")["rows"])
    assert any("قنا" in r["name"] for r in next(t for s in rm.sections for t in s.tables if t["key"] == "groups")["rows"])


def test_holders_are_admin_only(session, ds):
    rm, _ = service.build(session, ds, "en", admin=False)
    assert not any(t["key"] == "holders" for s in rm.sections for t in s.tables)


def test_pdf_and_excel_exports(session, ds):
    for lang in ("ar", "en"):
        rm, _ = service.build(session, ds, lang, admin=True)
        pdf = PdfExporter().render(rm)
        assert pdf[:5] == b"%PDF-" and len(pdf) > 20000
        wb = load_workbook(io.BytesIO(ExcelExporter().render(rm)))
        assert len(wb.sheetnames) >= 8
        cats = next(w for w in wb.worksheets if "Expenditure by category" in w.title or "حسب البند" in w.title)
        values = [c.value for row in cats.iter_rows(min_row=4) for c in row if isinstance(c.value, (int, float))]
        assert any(abs(v - 657500) < 0.01 for v in values)  # a real number, not text


def test_branch_names_inconsistent_period_is_excluded_from_branch_changes(session):
    d = service.ingest(session, monthly_branch_workbook(), "f2.xlsx", "tester")
    a = service.run_analysis(session, d)
    assert [k[1] for k in a["branch_name_inconsistent_periods"]] == [2]
    assert not any(o["rule"] == "mom_change" and o["subject_type"] == "branch" and o["period"] == (2026, 2) for o in a["outliers"])
    rm, _ = service.build(session, d, "en", admin=False)
    assert any("different branch names" in i["text"] for i in rm.sections[0].insights)


def test_display_taxonomy_groups_without_touching_originals(session, ds):
    set_setting(session, "custody.display_taxonomy", {"groups": {"علاقات عامة": ["Public Relation", "Public Relations"]}}, "admin")
    rm, a = service.build(session, ds, "ar", admin=False)
    cat_rows = next(t for s in rm.sections for t in s.tables if t["key"] == "categories")["rows"]
    assert {"Public Relation", "Public Relations"} <= {r["name"] for r in cat_rows}  # originals intact
    grouped = next(t for s in rm.sections for t in s.tables if t["key"] == "grouped")["rows"]
    assert next(r for r in grouped if r["group"] == "علاقات عامة")["total"] == Decimal("89000")
    assert any(g["basis"] in ("label_in_file", "similar_name") and "Public Relation" in g["members"] for g in ds.summary["suggested_groups"])


# ------------------------------------------------------------------ API
@pytest.fixture
def api(session, monkeypatch):
    monkeypatch.setenv("API_TOKENS", "tk-admin:admin:alice,tk-analyst:analyst:bob,tk-viewer:viewer:carol")
    get_settings.cache_clear()
    app.dependency_overrides[get_session] = lambda: (yield session)
    yield TestClient(app)
    app.dependency_overrides.clear()
    monkeypatch.delenv("API_TOKENS")
    get_settings.cache_clear()


def auth(role):
    return {"Authorization": f"Bearer tk-{role}"}


def test_api_upload_report_downloads_and_rbac(api):
    files = {"file": ("f3.xlsx", standard_gl())}
    assert api.post("/api/analysis/custody/datasets", files=files, headers=auth("viewer")).status_code == 403
    r = api.post("/api/analysis/custody/datasets", files=files, headers=auth("analyst"))
    assert r.status_code == 201
    did = r.json()["id"]
    assert r.json()["layout"] == "gl_settlement_lines" and r.json()["year"] == 2026
    assert api.post("/api/analysis/custody/datasets", files={"file": ("f3.xlsx", standard_gl())}, headers=auth("analyst")).status_code == 409
    junk = api.post("/api/analysis/custody/datasets", files={"file": ("x.xlsx", b"not excel")}, headers=auth("analyst"))
    assert junk.status_code == 400
    rep = api.get(f"/api/analysis/custody/datasets/{did}/report?lang=en", headers=auth("viewer")).json()
    assert rep["analysis"]["total"] == 871500
    assert not any(t["key"] == "holders" for s in rep["report"]["sections"] for t in s["tables"])  # viewer: no personal data
    adm = api.get(f"/api/analysis/custody/datasets/{did}/report?lang=en", headers=auth("admin")).json()
    assert any(t["key"] == "holders" for s in adm["report"]["sections"] for t in s["tables"])
    pdf = api.get(f"/api/analysis/custody/datasets/{did}/report.pdf?lang=ar", headers=auth("viewer"))
    assert pdf.status_code == 200 and pdf.content[:5] == b"%PDF-"
    xl = api.get(f"/api/analysis/custody/datasets/{did}/report.xlsx?lang=en", headers=auth("viewer"))
    assert xl.status_code == 200 and xl.content[:2] == b"PK"
    assert api.get(f"/api/analysis/custody/datasets/{did}/report?lang=xx", headers=auth("viewer")).status_code == 422
    assert api.delete(f"/api/analysis/custody/datasets/{did}", headers=auth("analyst")).status_code == 403


def test_api_thresholds_editable_without_code_change(api):
    files = {"file": ("f3.xlsx", standard_gl())}
    did = api.post("/api/analysis/custody/datasets", files=files, headers=auth("analyst")).json()["id"]
    base = api.get(f"/api/analysis/custody/datasets/{did}/report?lang=en", headers=auth("viewer")).json()["analysis"]["outliers"]
    assert api.put("/api/settings/custody.thresholds", json={"mom_abs": 10**9}, headers=auth("analyst")).status_code == 403
    assert api.put("/api/settings/custody.thresholds", json={"nope": 1}, headers=auth("admin")).status_code == 422
    assert api.put("/api/settings/custody.thresholds", json={"mom_abs": 10**9}, headers=auth("admin")).status_code == 200
    got = api.get("/api/settings/custody.thresholds", headers=auth("viewer")).json()
    assert got["effective"]["mom_abs"] == 10**9 and got["origin"]["mom_abs"] == "custom" and got["defaults"]["mom_abs"] != 10**9
    after = api.get(f"/api/analysis/custody/datasets/{did}/report?lang=en", headers=auth("viewer")).json()["analysis"]
    assert len(after["outliers"]) < len(base) and after["thresholds"]["mom_abs"] == 10**9
    # approving a display taxonomy: rejects a category listed in two groups
    bad = api.put("/api/settings/custody.display_taxonomy", json={"groups": {"a": ["X"], "b": ["X"]}}, headers=auth("admin"))
    assert bad.status_code == 422
