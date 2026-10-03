import io
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import get_settings
from app.core.analysis import AnalysisError
from app.core.reporting.excel import ExcelExporter
from app.core.reporting.pdf import PdfExporter
from app.db import get_session
from app.main import app
from app.modules.copier_analysis import engine, evidence, service
from app.modules.copier_analysis import invoice as inv_mod
from app.modules.copier_analysis import statement as st_mod
from app.modules.copier_analysis.adapter import CopierAdapter
from tests.copier_helpers import invoice_pdf, statement_workbook

TH = {"low_utilization_pct": 25, "very_low_utilization_pct": 5, "high_excess_pct": 50, "top_n": 10, "mom_change_pct": 25,
      "min_months_for_trend": 3, "reconcile_tolerance": 0.01}


# ---------------------------------------------------------------- statement
def test_statement_part1_is_primary_and_part2_is_matched_not_summed():
    st = st_mod.parse_statement_xlsx(statement_workbook())
    assert st.period == (2026, 7) and len(st.machines) == 5 and len(st.details) == 5
    assert [b.class_key for b in st.blocks] == ["mono_copier_5000", "printer_3000"]
    st_mod.validate_machines(st)
    st_mod.match_details(st, st.details, "xlsx_part2")
    assert {m.location_group for m in st.machines} == {"الشرقيه", "قنا"}   # governorate comes from Part 2 only
    codes = {i.code for i in st.issues}
    assert "branch_stray_marker" in codes and "block_total_missing" not in codes
    assert not any(c.startswith("machine_missing_in_detail") or c.startswith("detail_row_not_in_part1") for c in codes)
    # the stray marker is presentation only: the source text is kept
    assert any(m.branch_source.endswith("B") for m in st.machines) and st_mod.display_branch("فرع ألف B") == "فرع ألف"


def test_statement_disagreements_are_flagged_never_fixed():
    st = st_mod.parse_statement_xlsx(statement_workbook(drop_part2_row=True, break_total=True))
    st_mod.validate_machines(st)
    st_mod.match_details(st, st.details, "xlsx_part2")
    codes = {i.code for i in st.issues}
    assert "machine_missing_in_detail_xlsx_part2" in codes   # a Part 1 machine absent from Part 2 (like Abu Hummus)
    assert "block_total_differs" in codes                     # the block's stated total vs its rows
    miss = next(m for m in st.machines if m.branch_source == "فرع ألف" and m.class_key == "printer_3000")
    assert miss.location_group is None and miss.cons == 200   # Part 1 value untouched, location NOT invented


def test_classification_of_titles_and_machine_texts():
    assert st_mod.classify("زيارات شركه بدايتى 7- 2026   ماكينات الوان - 2000 نسخه")[0] == "color_copier_2000"
    assert st_mod.classify("ماكينات زيروكس الوان(A3) 2000نسخه")[0] == "color_a3_2000"
    assert st_mod.classify("طابعات - 3000 نسخه")[0] == "printer_3000"
    assert st_mod.classify("1000الوان")[0] == "color_copier_1000" and st_mod.classify("10000")[0] == "mono_copier_10000"
    assert st_mod.classify("بدون رقم") == (None, None)


# ---------------------------------------------------------------- invoice
def test_invoice_reader_rebuilds_lines_and_totals():
    inv = inv_mod.parse_invoice(invoice_pdf())
    assert inv.internal_no == "1976" and inv.electronic_id == "2BYHN2XAHNQ994SQTHQKJKZK10" and str(inv.issued_on) == "2026-08-09"
    assert [(ln.kind, ln.qty, ln.unit_price, ln.sales_total) for ln in inv.lines] == [
        ("rent", Decimal("3"), Decimal("1045"), Decimal("3135")), ("rent", Decimal("2"), Decimal("660"), Decimal("1320")),
        ("excess", Decimal("3000"), Decimal("0.319"), Decimal("957")), ("excess", Decimal("600"), Decimal("0.319"), Decimal("191.4"))]
    assert inv.lines[1].printers and inv.lines[2].packages == [5000]
    assert inv.totals["sales_total"] == Decimal("5603.4") and inv.period == (2026, 7)


def test_non_invoice_pdf_is_not_an_invoice():
    with pytest.raises(inv_mod.UnrecognisedInvoice):
        inv_mod.parse_invoice(b"%PDF-1.4\n%not an invoice\n")


# ---------------------------------------------------------------- engine
def _machines(**kw):
    st = st_mod.parse_statement_xlsx(statement_workbook(**kw))
    st_mod.validate_machines(st)
    st_mod.match_details(st, st.details, "xlsx_part2")
    return [{"source_ref": m.source_ref, "class_key": m.class_key, "package": m.package, "branch_key": m.branch_key,
             "branch_display": st_mod.display_branch(m.branch_source), "location_group": m.location_group, "group_kind": m.group_kind,
             "cons": m.cons, "exc": m.exc, "prev": m.prev, "cur": m.cur, "rent_detail": m.rent_detail, "flags": m.flags} for m in st.machines]


def _invoice(**kw):
    inv = inv_mod.parse_invoice(invoice_pdf(**kw))
    return {"totals": inv.totals, "lines": [{"line_no": ln.line_no, "description_short": ln.description[:30], "kind": ln.kind, "packages": ln.packages,
                                             "color": ln.color, "a3": ln.a3, "printers": ln.printers, "qty": ln.qty, "unit_price": ln.unit_price,
                                             "sales_total": ln.sales_total, "vat_rate": ln.vat_rate, "wht_rate": ln.wht_rate} for ln in inv.lines]}


def test_reconciliation_matches_and_costs_sum_to_the_invoice():
    a = engine.analyze_cycle(_machines(), _invoice(), TH)
    rec = a["reconciliation"]
    assert rec["summary"]["diff"] == 0 and rec["summary"]["ok"] >= 8 and not rec["unmapped_classes"]
    t = a["totals"]
    assert t["machines"] == 5 and t["consumption"] == 4000 + 8000 + 3500 + 200 + 3600
    assert t["excess_pages"] == 3600 and t["in_excess"] == 2
    assert t["rent"] == Decimal("3135") + Decimal("1320")                   # class rent × machines (invoice prices)
    assert t["excess_cost"] == Decimal("3600") * Decimal("0.319")
    assert t["cost"] == Decimal("5603.4")                                    # equals the invoice's total sales
    assert t["utilization"] == pytest.approx(19300 / 21000 * 100)


def test_reconciliation_flags_every_kind_of_difference():
    a = engine.analyze_cycle(_machines(), _invoice(qty_5000=4, exc_5000=3100), TH)
    ids = {c["id"]: c for c in a["reconciliation"]["checks"] if c["status"] == "diff"}
    assert ids["rent_qty_1"]["expected"] == 3 and ids["rent_qty_1"]["invoiced"] == 4          # one machine more on the invoice
    assert ids["excess_qty_3"]["diff"] == 100                                               # additional copies differ
    assert ids["total_expected_sales"]["status"] == "diff"
    b = engine.analyze_cycle(_machines(), _invoice(sales_wrong=True), TH)
    assert any(c["id"] == "total_line_sum" and c["status"] == "diff" for c in b["reconciliation"]["checks"])


def test_no_invoice_means_no_cost_and_no_reconciliation():
    a = engine.analyze_cycle(_machines(), None, TH)
    assert a["totals"]["cost"] == 0 and "cost" in a["unsupported"] and "reconciliation" not in a
    assert a["totals"]["cost_per_page"] is None or not a["has_invoice"]


def test_thresholds_drive_underutilisation():
    loose = engine.analyze_cycle(_machines(), _invoice(), {**TH, "low_utilization_pct": 90})
    strict = engine.analyze_cycle(_machines(), _invoice(), {**TH, "low_utilization_pct": 1})
    assert loose["bands"]["under_utilised"] > strict["bands"]["under_utilised"]
    assert [r["branch_display"] for r in strict["under_utilised"]] == []


# ---------------------------------------------------------------- evidence = supporting only
def test_evidence_matching_statuses():
    res = evidence.match_pages(
        [{"page_no": 1, "counter_value": 5000}, {"page_no": 2, "counter_value": 5000}, {"page_no": 3, "counter_value": 12345},
         {"page_no": 4, "counter_value": None}], [{"source_ref": "r1", "cur": 5000}])
    assert [r["status"] for r in res] == ["matched", "duplicate", "no_match", "unreadable"]
    assert evidence.read_counter("Counters Printed Pages 144931\nx")[0] == 144931
    assert evidence.printed_at("a\nb 26/07/2026 10:44\nc 09/12/2025 11:21:01")[:10] == "26/07/2026"


# ---------------------------------------------------------------- cycle / adapter / API
@pytest.fixture
def adapter():
    return CopierAdapter()


def _up(adapter, session, content, name):
    return adapter.ingest(session, content, name, "tester", {})


def test_files_of_one_month_form_one_cycle_in_any_order(adapter, session):
    r1 = _up(adapter, session, invoice_pdf(), "invoice.pdf")
    r2 = _up(adapter, session, statement_workbook(), "statement.xlsx")
    assert r1.item_id == r2.item_id and r2.meta["has_statement"] and r2.meta["has_invoice"]
    with pytest.raises(AnalysisError) as e:
        _up(adapter, session, statement_workbook(), "again.xlsx")
    assert e.value.status == 409
    with pytest.raises(AnalysisError) as e:
        _up(adapter, session, b"not an excel", "x.xlsx")
    assert e.value.status == 400


def test_cycle_report_is_computed_honest_and_exportable(adapter, session):
    _up(adapter, session, statement_workbook(), "statement.xlsx")
    _up(adapter, session, invoice_pdf(), "invoice.pdf")
    for lang in ("ar", "en"):
        rm, a = adapter.build(session, 1, lang, True, {})
        keys = {t["key"] for s in rm.sections for t in s.tables}
        assert {"classes", "under", "exceeding", "branches", "locations", "recon", "invoice", "issues", "settings"} <= keys
        assert rm.meta["filters"]["dimensions"][0]["key"] == "governorates"
        assert any("5,603" in i["text"] or "5603" in i["text"] for i in rm.meta["summary_items"])
        assert PdfExporter().render(rm)[:5] == b"%PDF-"
        wb = load_workbook(io.BytesIO(ExcelExporter().render(rm)))
        assert len(wb.sheetnames) >= 8


def test_filters_recompute_and_hide_reconciliation(adapter, session):
    _up(adapter, session, statement_workbook(), "statement.xlsx")
    _up(adapter, session, invoice_pdf(), "invoice.pdf")
    rm, a = adapter.build(session, 1, "en", True, {"governorates": ["قنا"]})
    assert a["totals"]["machines"] == 1 and "reconciliation" not in a
    assert not any(t["key"] == "recon" for s in rm.sections for t in s.tables)
    rm2, a2 = adapter.build(session, 1, "en", True, {"classes": ["printer_3000"]})
    assert a2["totals"]["machines"] == 2 and a2["totals"]["cost"] == Decimal("1320") + Decimal("600") * Decimal("0.319")


def test_evidence_never_changes_kpis_or_costs(adapter, session):
    from app.models import AnalysisDataset, CopierEvidence
    _up(adapter, session, statement_workbook(), "statement.xlsx")
    _up(adapter, session, invoice_pdf(), "invoice.pdf")
    _, before = adapter.build(session, 1, "en", True, {})
    ds = AnalysisDataset(module_id="copier_analysis", layout="copier_evidence", scope_label="evidence_only", file_name="scan.pdf", file_hash="x" * 64,
                         cycle_id=1, role="evidence", status="ready", facts_count=2)
    session.add(ds)
    session.flush()
    session.add_all([CopierEvidence(dataset_id=ds.id, page_no=1, counter_value=999999999, status="no_match", note="conflict"),
                     CopierEvidence(dataset_id=ds.id, page_no=2, counter_value=None, status="unreadable")])
    session.commit()
    rm, after = adapter.build(session, 1, "en", True, {})
    for k in ("machines", "consumption", "excess_pages", "rent", "excess_cost", "cost"):
        assert before["totals"][k] == after["totals"][k]
    ev = next(s for s in rm.sections if s.key == "evidence")
    assert "supporting evidence" in ev.insights[0]["text"] and any(t["key"] == "evidence" for t in ev.tables)


def test_trend_over_two_months_and_reading_continuity(adapter, session):
    _up(adapter, session, statement_workbook(month=7), "july.xlsx")
    _up(adapter, session, statement_workbook(month=8, shift=20000), "august.xlsx")   # prev readings = July's current ones? (shifted)
    items = adapter.list_items(session)
    assert items[0]["id"] == "all" and len(items) == 3
    rm, a = adapter.build(session, "all", "en", True, {})
    assert [r["period"] for r in a["trend"]["rows"]] == [(2026, 7), (2026, 8)]
    assert a["trend"]["rows"][1]["consumption_mom"] == 0 and a["trend"]["indicative"] is True
    assert a["trend"]["continuity"][0]["machines"] == 5
    assert any(t["key"] == "trend" for s in rm.sections for t in s.tables)


# ---------------------------------------------------------------- API
@pytest.fixture
def api(session, monkeypatch):
    monkeypatch.setenv("API_TOKENS", "tk-admin:admin:alice,tk-analyst:analyst:bob,tk-viewer:viewer:carol")
    get_settings.cache_clear()
    app.dependency_overrides[get_session] = lambda: (yield session)
    yield TestClient(app)
    app.dependency_overrides.clear()
    monkeypatch.delenv("API_TOKENS")
    get_settings.cache_clear()


def auth(role):
    return {"Authorization": f"Bearer tk-{role}"}


def test_generic_api_serves_both_modules_and_exports(api):
    mods = api.get("/api/analysis", headers=auth("viewer")).json()
    assert {m["key"] for m in mods} == {"custody", "copiers", "aramex", "procurement"}
    base = "/api/analysis/copiers/datasets"
    assert api.post(base, files={"file": ("s.xlsx", statement_workbook())}, headers=auth("viewer")).status_code == 403
    r = api.post(base, files={"file": ("s.xlsx", statement_workbook())}, headers=auth("analyst"))
    assert r.status_code == 201 and r.json()["uploaded"]["role"] == "statement"
    r = api.post(base, files={"file": ("i.pdf", invoice_pdf())}, headers=auth("analyst"))
    assert r.status_code == 201 and r.json()["has_invoice"]
    cid = r.json()["id"]
    assert api.post(base, files={"file": ("i.pdf", invoice_pdf())}, headers=auth("analyst")).status_code == 409
    rep = api.get(f"{base}/{cid}/report?lang=en", headers=auth("viewer")).json()
    assert rep["analysis"]["reconciliation"]["summary"]["diff"] == 0
    part = api.get(f"{base}/{cid}/report?lang=en&class=printer_3000", headers=auth("viewer")).json()
    assert part["analysis"]["totals"]["machines"] == 2
    pdf = api.get(f"{base}/{cid}/report.pdf?lang=ar&governorate=%D9%82%D9%86%D8%A7", headers=auth("viewer"))
    assert pdf.status_code == 200 and pdf.content[:5] == b"%PDF-"
    assert api.get(f"{base}/{cid}/report.xlsx?lang=en", headers=auth("viewer")).content[:2] == b"PK"
    assert api.get(f"{base}/999/report", headers=auth("viewer")).status_code == 404
    assert api.get("/api/analysis/nope/datasets", headers=auth("viewer")).status_code == 404
    # custody keeps working on the same generic routes
    assert api.get("/api/analysis/custody/datasets", headers=auth("viewer")).status_code == 200


def test_copier_thresholds_are_editable_settings(api):
    base = "/api/analysis/copiers/datasets"
    api.post(base, files={"file": ("s.xlsx", statement_workbook())}, headers=auth("analyst"))
    got = api.get("/api/settings/copier.thresholds", headers=auth("viewer")).json()
    assert got["origin"]["low_utilization_pct"] == "default" and got["defaults"]["low_utilization_pct"] == 25
    assert api.put("/api/settings/copier.thresholds", json={"low_utilization_pct": 90}, headers=auth("analyst")).status_code == 403
    assert api.put("/api/settings/copier.thresholds", json={"nope": 1}, headers=auth("admin")).status_code == 422
    assert api.put("/api/settings/copier.thresholds", json={"low_utilization_pct": 90}, headers=auth("admin")).status_code == 200
    rep = api.get(f"{base}/1/report?lang=en", headers=auth("viewer")).json()
    assert rep["analysis"]["thresholds"]["low_utilization_pct"] == 90 and rep["analysis"]["thresholds_origin"]["low_utilization_pct"] == "custom"
    assert api.get("/api/settings/other.thresholds", headers=auth("viewer")).status_code == 404




# ---------------------------------------------------------------- scanned status pages (needs tesseract)
def _scan_pdf(counters: list[str]) -> bytes:
    from PIL import Image, ImageDraw, ImageFont
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 34)
    pages = []
    for c in counters:
        im = Image.new("L", (1240, 700), 255)
        d = ImageDraw.Draw(im)
        d.text((60, 40), "Status Page   Printer P-4030DN   26/07/2026 10:44", font=font, fill=0)
        d.text((60, 200), "Sleep Timer: 120 Minutes     Counters", font=font, fill=0)
        d.text((60, 260), f"Printed Pages {c}", font=font, fill=0)
        pages.append(im)
    buf = io.BytesIO()
    pages[0].save(buf, format="PDF", save_all=True, append_images=pages[1:], resolution=200)
    return buf.getvalue()


@pytest.mark.skipif(__import__("shutil").which("tesseract") is None or __import__("shutil").which("pdftoppm") is None, reason="OCR tools missing")
def test_scanned_status_pages_are_matched_as_evidence_only(adapter, session):
    from sqlalchemy import select

    from app.models import AnalysisDataset
    _up(adapter, session, statement_workbook(), "statement.xlsx")
    _up(adapter, session, invoice_pdf(), "invoice.pdf")
    _, before = adapter.build(session, 1, "en", True, {})
    res = _up(adapter, session, _scan_pdf(["4500", "123456"]), "scan.pdf")      # 4500 = a current reading of the statement
    assert res.status == 202 and res.meta["uploaded"]["role"] == "evidence" and res.background is not None
    ds = session.scalar(select(AnalysisDataset).where(AnalysisDataset.role == "evidence"))
    service.process_evidence(session, ds.id)
    rm, after = adapter.build(session, 1, "en", True, {})
    ev = next(t for s in rm.sections for t in s.tables if t["key"] == "evidence")["rows"]
    assert [r["status"] for r in ev] == ["Matches a statement reading", "No reading matches (review)"]
    assert before["totals"] == after["totals"]                      # evidence feeds no KPI/cost/total
    assert any("1 of 2" in i["text"] for i in next(s for s in rm.sections if s.key == "quality").insights)
